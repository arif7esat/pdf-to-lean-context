#!/usr/bin/env python3
"""
slayt_paketle.py — Bir klasördeki tüm slaytları (PDF / PPT / PPTX / ODP, karışık)
tek seferde "Markdown + slayt görselleri + README + kalite raporu" paketlerine çevirir.

İlke: Dönüşümü dil modeli değil, bu betik yapar. Metin yalnızca dosyada gerçekten
yazanlardan alınır; şemalar metne "yorumlanmaz", görsel olarak bırakılıp işaretlenir.
Kalite raporu "doğrulandı" demez; ölçüm verir ve kontrol edilecek slaytları listeler.

COLAB'DA (tek paylaşım linki, drive.mount GEREKMEZ)
    Bu dosyanın tamamını bir Colab hücresine yapıştırıp çalıştır; ses ve kitap
    betiklerinde kullandığın AYNI ders klasörü linkini yapıştır.
        <Ders>/DersKaynaklari/Slaytlar/         <- slayt PDF/PPT/PPTX (alt klasör serbest)
        <Ders>/DersKaynaklari/Slayt_Paketleri/  <- çıktı (betik oluşturur)
        <Ders>/DersKaynaklari/Kitap/*_Bolumler/ <- varsa otomatik --sozluk olarak kullanılır
    YALNIZCA 'Slaytlar' okunur, YALNIZCA 'Slayt_Paketleri' içine yazılır.

KULLANIM (lokal)
    python slayt_paketle.py GIRIS_KLASORU CIKIS_KLASORU [--dpi 110] [--yeniden] [--sozluk KITAP_MD_KLASORU]

    GIRIS_KLASORU : slayt dosyalarının olduğu klasör (alt klasörler de taranır)
    CIKIS_KLASORU : paketlerin yazılacağı klasör
    --dpi         : slayt görsellerinin çözünürlüğü (varsayılan 110)
    --yeniden     : daha önce işlenmiş ve değişmemiş dosyaları da yeniden işle
    --sozluk      : (isteğe bağlı) kitabın Markdown klasörü; "weakne ss" gibi bölünmüş kelimeleri
                    onarmak için ek sözlük olarak kullanılır

KURULUM
    pip install pymupdf python-pptx pytesseract
    OCR (isteğe bağlı, resme dönüşmüş tablo/metinleri okur): Tesseract
      Colab : betik kendisi kurar          macOS : brew install tesseract tesseract-lang
    Colab'da .ppt/.pptx/.odp dosyaları Google Slaytlar ile çevrilir (Drive izni verilirse);
    Google başarısız olursa ya da lokalde çalıştırılırsa LibreOffice kullanılır:
      Colab/Linux : apt-get install -y libreoffice-impress
      macOS       : brew install --cask libreoffice

DOSYA EŞLEŞTİRME
    - Aynı adlı .pptx ve .pdf varsa: metin .pptx'ten, görseller .pdf'ten alınır (en iyisi).
    - Yalnızca .pptx / .ppt / .odp varsa: metin sunudan, görseller LibreOffice ile üretilir.
    - Yalnızca .pdf varsa: metin ve görseller PDF'ten alınır.
"""

import argparse
import hashlib
import json
import re
import shutil
import subprocess
import sys
import tempfile
from collections import Counter
from pathlib import Path

COLAB = "google.colab" in sys.modules or Path("/content").exists()

try:
    import pymupdf
except ImportError:
    try:
        import fitz as pymupdf  # eski sürümler
    except ImportError:
        if not COLAB:
            raise
        subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", "pymupdf", "python-pptx"])
        import pymupdf

SURUM = "1.9"
SAYFA_GORSELLERI = "gorselli"   # "gorselli": yalnızca resim/şema/grafik içeren slaytlar · "hepsi": tüm slaytlar
OCR = True                  # görselli slaytlarda resme dönüşmüş yazıyı Tesseract ile oku
OCR_DILLERI = "eng+tur"     # kurulu olmayan dil otomatik atlanır
OCR_GUVEN_ESIGI = 90        # ortalama güveni bunun altındaki okuma md'ye eklenmez (ölçüm: temiz okumalar
                            # %95-96, bozuk/düşük çözünürlüklü resimlerdeki çöp okumalar %72-88)
OCR_SOZLUK_ORANI = 0.9      # sözlük büyükse okunan kelimelerin en az bu kadarı sözlükte olmalı
JSON_KLASORU = "Jsonlar"   # kalite raporları paketlerin dışında, çıktı kökündeki bu klasörde toplanır
EK_SOZLUK = None  # --sozluk ile verilen klasördeki .md/.txt dosyalarının kelimeleri
SUNU_UZANTILARI = {".pptx", ".ppt", ".odp", ".pps", ".ppsx"}
BULLET_KARAKTERLERI = set("•●○◦▪■□►▸▶➢➤✓✔❖◆◇–—-*·")
SEMBOL_FONTLARI = ("wingdings", "symbol", "zapfdingbats", "webdings", "dingbat")


# ──────────────────────────────────────────────────────────────────────────────
# Genel yardımcılar
# ──────────────────────────────────────────────────────────────────────────────

def sha1(yol: Path) -> str:
    h = hashlib.sha1()
    with open(yol, "rb") as f:
        for parca in iter(lambda: f.read(1 << 20), b""):
            h.update(parca)
    return h.hexdigest()


def md5(yol: Path) -> str:
    h = hashlib.md5()
    with open(yol, "rb") as f:
        for parca in iter(lambda: f.read(1 << 20), b""):
            h.update(parca)
    return h.hexdigest()


def kelimeler(metin: str) -> list:
    """Karşılaştırma için normalleştirilmiş kelime listesi (küçük harf, yalnız harf/rakam)."""
    return re.findall(r"[^\W_]+", metin.lower())


def kapsama(kaynak: str, cikti: str):
    """Kaynaktaki kelimelerin çıktıda bulunma oranı + eksik/fazla kelimeler (çoklu küme)."""
    k, c = Counter(kelimeler(kaynak)), Counter(kelimeler(cikti))
    toplam = sum(k.values())
    ortak = sum((k & c).values())
    eksik = list((k - c).elements())
    fazla = list((c - k).elements())
    oran = 1.0 if toplam == 0 else ortak / toplam
    return oran, eksik, fazla, toplam


def satir_birlestir(onceki: str, sonraki: str) -> str:
    """Satır sonunda kelime içi tire varsa tireyi koruyarak bitişik birleştir."""
    onceki, sonraki = onceki.rstrip(), sonraki.strip()
    if not onceki:
        return sonraki
    if len(onceki) >= 2 and onceki.endswith("-") and onceki[-2].isalpha() and sonraki[:1].isalpha():
        return onceki + sonraki          # cost- + effective -> cost-effective (tire korunur)
    return onceki + " " + sonraki


def md_kacis(metin: str) -> str:
    metin = metin.strip()
    if re.match(r"^(#|>|\d+[.)]\s|[-+*]\s)", metin):
        metin = "\\" + metin
    return metin


def hucre(metin: str) -> str:
    return metin.replace("|", "\\|").strip()


def kume(degerler, tolerans):
    """Sayıları toleransa göre kümele; küme merkezlerini ve üye sayılarını döndür."""
    kumeler = []
    for d in sorted(degerler):
        if kumeler and d - kumeler[-1][-1] <= tolerans:
            kumeler[-1].append(d)
        else:
            kumeler.append([d])
    return [(sum(k) / len(k), len(k)) for k in kumeler]


def soffice_bul():
    for ad in ("soffice", "libreoffice"):
        yol = shutil.which(ad)
        if yol:
            return yol
    mac = Path("/Applications/LibreOffice.app/Contents/MacOS/soffice")
    return str(mac) if mac.exists() else None


HARICI_DONUSTURUCU = None   # Colab'da Google Slaytlar dönüştürücüsü atanır: f(kaynak, hedef_tur, cikis) -> Path
LIBREOFFICE_KURUCU = None   # Colab'da LibreOffice'i gerektiğinde kuran fonksiyon


def sunu_cevir(kaynak: Path, hedef_tur: str, cikis: Path):
    """Sunuyu pptx/pdf'e çevirir; (dosya, kullanılan araç) döndürür.
    Önce varsa Google Slaytlar, olmazsa LibreOffice kullanılır."""
    if HARICI_DONUSTURUCU:
        try:
            return HARICI_DONUSTURUCU(kaynak, hedef_tur, cikis), "Google Slaytlar"
        except Exception as e:
            print(f"   ⚠ Google Slaytlar dönüştürmesi başarısız ({e}); LibreOffice deneniyor.")
    return soffice_cevir(kaynak, hedef_tur, cikis), "LibreOffice"


def soffice_cevir(kaynak: Path, hedef_tur: str, cikis: Path) -> Path:
    soffice = soffice_bul()
    if not soffice and LIBREOFFICE_KURUCU:
        LIBREOFFICE_KURUCU()
        soffice = soffice_bul()
    if not soffice:
        raise RuntimeError("LibreOffice bulunamadı (sunu dosyaları için gerekli). KURULUM bölümüne bakın.")
    cikis.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory() as profil:
        komut = [soffice, f"-env:UserInstallation=file://{profil}", "--headless",
                 "--convert-to", hedef_tur, "--outdir", str(cikis), str(kaynak)]
        subprocess.run(komut, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, timeout=300)
    sonuc = cikis / (kaynak.stem + "." + hedef_tur)
    if not sonuc.exists():
        raise RuntimeError(f"LibreOffice dönüştürmesi başarısız: {kaynak.name} -> {hedef_tur}")
    return sonuc


# ──────────────────────────────────────────────────────────────────────────────
# PDF'ten satır çıkarma
# ──────────────────────────────────────────────────────────────────────────────

def bullet_mi(metin: str, font: str) -> bool:
    t = metin.strip()
    if len(t) != 1:
        return False
    f = font.lower()
    if any(s in f for s in SEMBOL_FONTLARI):
        return True
    if 0xF000 <= ord(t) <= 0xF0FF:  # sembol fontlarının özel kullanım alanı
        return True
    if t in BULLET_KARAKTERLERI:
        return True
    return t == "o" and "courier" in f


def sayfa_satirlari(sayfa, duzeltmeler: list):
    """Sayfadaki metin satırlarını ve madde işaretlerini karakter düzeyinde çıkarır.

    Ölçülebilir biçimde yanlış olan boşluklar (normal boşluktan belirgin dar,
    iki harfin arasında) kaldırılır ve her biri `duzeltmeler` listesine yazılır.
    """
    satirlar, isaretler = [], []
    for blok in sayfa.get_text("rawdict")["blocks"]:
        if blok["type"] != 0:
            continue
        for satir in blok["lines"]:
            spanlar = [s for s in satir["spans"] if "".join(c["c"] for c in s["chars"]).strip() != ""
                       or len(s["chars"]) > 0]
            # Satır başındaki madde işaretini ayır
            while spanlar:
                ilk = spanlar[0]
                t = "".join(c["c"] for c in ilk["chars"])
                if bullet_mi(t, ilk["font"]) and len(spanlar) >= 1:
                    b = ilk["bbox"]
                    isaretler.append({"x": b[0], "yc": (b[1] + b[3]) / 2, "y0": b[1], "y1": b[3]})
                    spanlar = spanlar[1:]
                    if spanlar:
                        continue
                break
            if not spanlar:
                continue
            karakterler = []  # (char, bbox, size, font, span_index)
            for si, s in enumerate(spanlar):
                for c in s["chars"]:
                    karakterler.append((c["c"], c["bbox"], s["size"], s["font"], si, len(s["chars"])))
            if not "".join(k[0] for k in karakterler).strip():
                continue
            # Normal boşluk genişliği (bu satırdaki)
            bosluk_oranlari = [(k[1][2] - k[1][0]) / k[2] for k in karakterler if k[0] == " " and k[2] > 0]
            normal = sorted(bosluk_oranlari)[len(bosluk_oranlari) // 2] if bosluk_oranlari else 0.25
            metin = []
            for i, (ch, bb, boyut, font, si, span_uzunluk) in enumerate(karakterler):
                if ch == " " and 0 < i < len(karakterler) - 1:
                    sol, sag = karakterler[i - 1][0], karakterler[i + 1][0]
                    oran = (bb[2] - bb[0]) / boyut if boyut else normal
                    tek_bosluk_span = span_uzunluk == 1
                    if sol.isalpha() and sag.isalpha() and sol.islower() and sag.islower() and (
                            oran < 0.72 * normal or tek_bosluk_span):
                        once = "".join(k[0] for k in karakterler[max(0, i - 12):i]).split(" ")[-1]
                        sonra = "".join(k[0] for k in karakterler[i + 1:i + 13]).split(" ")[0]
                        duzeltmeler.append(f"'{once} {sonra}' → '{once}{sonra}'")
                        continue
                metin.append(ch)
            metin = "".join(metin)
            metin = re.sub(r"\s+", " ", metin).strip()
            if not metin:
                continue
            x0 = min(k[1][0] for k in karakterler if k[0].strip())
            y0 = min(k[1][1] for k in karakterler)
            x1 = max(k[1][2] for k in karakterler)
            y1 = max(k[1][3] for k in karakterler)
            boyut = max(k[2] for k in karakterler)
            satirlar.append({"x0": x0, "y0": y0, "x1": x1, "y1": y1, "boyut": boyut,
                             "metin": metin, "bullet_x": None})
    return satirlar, isaretler


def isaretleri_eslestir(satirlar, isaretler):
    """Her madde işaretini sağındaki, dikeyde örtüşen satıra bağlar."""
    sahipsiz = 0
    for isaret in isaretler:
        adaylar = []
        for s in satirlar:
            yukseklik = s["y1"] - s["y0"]
            if s["x0"] > isaret["x"] and s["x0"] - isaret["x"] < 90 and \
                    s["y0"] - 0.3 * yukseklik <= isaret["yc"] <= s["y1"] + 0.3 * yukseklik and s["bullet_x"] is None:
                adaylar.append((s["x0"] - isaret["x"], abs((s["y0"] + s["y1"]) / 2 - isaret["yc"]), id(s), s))
        if adaylar:
            adaylar.sort(key=lambda a: (a[1] > 0.6 * (a[3]["y1"] - a[3]["y0"]), a[0], a[1]))
            adaylar[0][3]["bullet_x"] = isaret["x"]
        else:
            sahipsiz += 1
    return sahipsiz


def sozluk_birlestir(satirlar, duzeltme_listeleri, ek_sozluk=None):
    """Kelime içine düşmüş sahte boşlukları destenin kendi kelime dağarcığıyla onarır.

    "haza rd" gibi bir çift, yalnızca (1) bitişik hâli ("hazard") destede başka bir yerde
    geçiyorsa ve (2) parçalardan en az biri destede hiçbir yerde tek başına kelime olarak
    geçmiyorsa birleştirilir. "in to" gibi iki gerçek kelime asla birleştirilmez.
    Her düzeltme raporlanır.
    """
    desen = re.compile(r"\b([A-Za-z]+) (?=([a-z]+)\b)")   # örtüşen çiftler için ileri bakış
    sayac = Counter(k.lower() for s in satirlar for k in re.findall(r"[A-Za-z]+", s["metin"]))
    aday_sol, aday_sag = Counter(), Counter()

    def bilinen(kelime):
        return kelime in sayac or (ek_sozluk is not None and kelime in ek_sozluk)

    for s in satirlar:
        for m in desen.finditer(s["metin"]):
            if bilinen((m.group(1) + m.group(2)).lower()):
                aday_sol[m.group(1).lower()] += 1
                aday_sag[m.group(2).lower()] += 1

    def yalniz_parca(p):
        return sayac[p] == aday_sol[p] + aday_sag[p]

    for s in satirlar:
        degisti = True
        while degisti:
            degisti = False
            for m in desen.finditer(s["metin"]):
                sol, sag = m.group(1), m.group(2)
                if bilinen((sol + sag).lower()) and (yalniz_parca(sol.lower()) or yalniz_parca(sag)):
                    duzeltme_listeleri[id(s)].append(f"'{sol} {sag}' → '{sol}{sag}' (sözlük)")
                    s["metin"] = s["metin"][:m.start()] + sol + s["metin"][m.end():]
                    degisti = True
                    break


def normal_metin(metin: str) -> str:
    return re.sub(r"\s+", " ", re.sub(r"\d+", "#", metin.strip().lower()))


# ──────────────────────────────────────────────────────────────────────────────
# PDF sayfasını yapılandırma (başlık / madde / tablo / şekil)
# ──────────────────────────────────────────────────────────────────────────────

def sekil_bolgeleri(sayfa):
    """Arka plan dışındaki raster görseller ve vektör çizim kümeleri."""
    A = sayfa.rect.get_area()
    W = sayfa.rect.width
    rasterlar, vektor = [], []
    for blok in sayfa.get_text("dict")["blocks"]:
        if blok["type"] == 1:
            r = pymupdf.Rect(blok["bbox"])
            oran = r.get_area() / A
            if 0.02 <= oran < 0.9:
                rasterlar.append(r)
    for d in sayfa.get_drawings():
        r = d["rect"]
        if r.get_area() / A > 0.8:
            continue                                   # tam sayfa arka plan
        if (r.height < 4 and r.width > 0.4 * W) or (r.width < 4 and r.height < 4):
            continue                                   # başlık çizgisi, tablo çizgisi, nokta
        if r.height < 4 or r.width < 4:
            vektor.append(("cizgi", r))
            continue
        vektor.append(("sekil", r))
    return rasterlar, vektor


def _ayni_hiza(a, b):
    """İki satır aynı yatay hizada mı? Dikey örtüşme, kısa olanın yüksekliğinin yarısından fazlaysa evet.
    (OCR'da kalın/uzun harfli kelimelerin kutusu şişebildiği için üst kenar farkı tek başına güvenilmez.)"""
    ortusme = min(a["y1"], b["y1"]) - max(a["y0"], b["y0"])
    return ortusme > 0.5 * max(1, min(a["y1"] - a["y0"], b["y1"] - b["y0"]))


def tablo_bul(govde):
    """Aynı yükseklikte, ayrık sütunlarda duran satırlardan tablo çıkarır. Yoksa None."""
    satir_gruplari = []
    for s in sorted(govde, key=lambda s: (s["y0"], s["x0"])):
        if s["bullet_x"] is not None:
            continue
        ilk = satir_gruplari[-1][0] if satir_gruplari else None
        if ilk and (abs(ilk["y0"] - s["y0"]) < 3 or _ayni_hiza(ilk, s)):
            satir_gruplari[-1].append(s)
        else:
            satir_gruplari.append([s])
    for g in satir_gruplari:          # aynı hizadaki parçalar soldan sağa (y'si 1 birim önde olan sağ sütun olabilir)
        g.sort(key=lambda s: s["x0"])
    cok_sutunlu = [g for g in satir_gruplari if len(g) >= 2 and
                   all(g[i + 1]["x0"] - g[i]["x1"] > 8 for i in range(len(g) - 1))]
    if len(cok_sutunlu) < 3:
        return None
    sutunlar = [m for m, _ in kume([s["x0"] for g in cok_sutunlu for s in g], 15)]
    if len(sutunlar) < 2:
        return None
    ust = min(s["y0"] for s in cok_sutunlu[0]) - 2
    tablo_satirlari = [s for s in govde if s["y0"] >= ust and s["bullet_x"] is None and
                       min(abs(s["x0"] - c) for c in sutunlar) < 15]
    if not tablo_satirlari:
        return None
    tablo_satirlari.sort(key=lambda s: (s["y0"], s["x0"]))
    # Görsel satırlar (aynı hizadaki parçalar tek satır) ve aralarındaki adım
    hizalar = []
    for s in tablo_satirlari:
        if hizalar and (abs(hizalar[-1][0]["y0"] - s["y0"]) < 3 or _ayni_hiza(hizalar[-1][0], s)):
            hizalar[-1].append(s)
        else:
            hizalar.append([s])
    adimlar = [hizalar[i + 1][0]["y0"] - hizalar[i][0]["y0"] for i in range(len(hizalar) - 1)]
    temel = sorted(adimlar)[len(adimlar) // 4] if adimlar else 0

    def ilk_sutunda(hiza):
        return any(min(range(len(sutunlar)), key=lambda k: abs(p["x0"] - sutunlar[k])) == 0 for p in hiza)
    ilkler = [ilk_sutunda(h) for h in hizalar]
    tekduze = all(ilkler) and bool(adimlar) and max(adimlar) < 1.3 * temel
    # Satır sınırı kuralları (ölçülmüş: satır aralığı tek başına OCR oynaklığında güvenilmez):
    # - ilk sütunda (terim sütunu) yazı yoksa hiza önceki satırın devamıdır;
    # - ilk sütunda yazı varsa ve önceki hizanın ilk sütunu boşsa yeni satırdır;
    # - ardışık iki hizanın ikisinde de ilk sütun doluysa (çok satırlı terim ya da tek satırlık satırlar)
    #   karar aralığa göre verilir; tüm satırlar tek satırlıksa her hiza yeni satırdır.
    satirlar = []
    for i, hiza in enumerate(hizalar):
        if i == 0:
            yeni = True
        elif not ilkler[i]:
            yeni = False
        elif tekduze or not ilkler[i - 1]:
            yeni = True
        else:
            yeni = adimlar[i - 1] > 1.3 * temel
        if yeni:
            satirlar.append([""] * len(sutunlar))
        for p in hiza:
            sutun = min(range(len(sutunlar)), key=lambda k: abs(p["x0"] - sutunlar[k]))
            satirlar[-1][sutun] = satir_birlestir(satirlar[-1][sutun], p["metin"])
    # İlk hücresi boş satır, önceki satırın devamıdır (hücre içi satır sonu)
    birlesik = []
    for r in satirlar:
        if birlesik and not r[0].strip():
            birlesik[-1] = [satir_birlestir(a, b) if b else a for a, b in zip(birlesik[-1], r)]
        else:
            birlesik.append(r)
    if len(birlesik) < 3:          # şema etiketlerinden sahte tablo çıkmasın
        return None
    return {"y": ust, "satirlar": birlesik, "kaynak": tablo_satirlari, "sutunlar": sutunlar}


def etiketleri_grupla(satirlar):
    """Şekil içindeki, üst üste duran ve ortalanmış satırları tek etikette birleştirir."""
    satirlar = sorted(satirlar, key=lambda s: (s["y0"], s["x0"]))
    gruplar = []
    for s in satirlar:
        merkez = (s["x0"] + s["x1"]) / 2
        for g in gruplar:
            son = g[-1]
            if abs((son["x0"] + son["x1"]) / 2 - merkez) < 20 and 0 <= s["y0"] - son["y1"] < 0.7 * s["boyut"]:
                g.append(s)
                break
        else:
            gruplar.append([s])
    # Aynı yükseklikte yan yana duran öbekler soldan sağa, sonra yukarıdan aşağı sıralanır
    gruplar.sort(key=lambda g: g[0]["y0"])
    siralar = []
    for g in gruplar:
        yuk = g[0]["y1"] - g[0]["y0"]
        if siralar and abs(g[0]["y0"] - siralar[-1][0][0]["y0"]) < 0.6 * max(yuk, 1):
            siralar[-1].append(g)
        else:
            siralar.append([g])
    gruplar = [g for sira in siralar for g in sorted(sira, key=lambda g: g[0]["x0"])]
    etiketler = []
    for g in gruplar:
        metin = ""
        for s in g:
            metin = satir_birlestir(metin, s["metin"])
        etiketler.append(metin)
    return etiketler


def pdf_sayfasi_yapilandir(sayfa, satirlar, bullet_seviyeleri):
    H = sayfa.rect.height
    sonuc = {"baslik": "", "bloklar": [], "sekil_etiketleri": [], "bayraklar": []}
    rasterlar, vektor = sekil_bolgeleri(sayfa)
    R = sayfa.rect
    sonuc["gorsel_alanlari"] = [(r.x0 / R.width, r.y0 / R.height, r.x1 / R.width, r.y1 / R.height) for r in rasterlar]
    if not satirlar:
        if rasterlar:
            sonuc["bayraklar"].append("görsel/şema (metni PDF'te yok, görsele bak)")
        return sonuc

    # --- Başlık: sayfanın üst kısmındaki en büyük yazı
    ust_bolge = [s for s in satirlar if s["y0"] < 0.28 * H and s["bullet_x"] is None]
    govde_boyutlari = sorted(s["boyut"] for s in satirlar)
    medyan = govde_boyutlari[len(govde_boyutlari) // 2]
    baslik_satirlari = []
    if ust_bolge:
        enbuyuk = max(s["boyut"] for s in ust_bolge)
        if enbuyuk >= 1.15 * medyan or len(satirlar) <= 2 or enbuyuk >= 30:
            baslik_satirlari = [s for s in ust_bolge if s["boyut"] >= enbuyuk - 1]
    if not baslik_satirlari:
        enbuyuk = max(s["boyut"] for s in satirlar)
        adaylar = [s for s in satirlar if s["boyut"] >= enbuyuk - 1 and s["bullet_x"] is None]
        digerleri = [s["boyut"] for s in satirlar if s not in adaylar]
        if adaylar and (not digerleri or enbuyuk >= 1.3 * max(digerleri)):
            baslik_satirlari = adaylar
    baslik_satirlari.sort(key=lambda s: (s["y0"], s["x0"]))
    for s in baslik_satirlari:
        sonuc["baslik"] = satir_birlestir(sonuc["baslik"], s["metin"])
    if not sonuc["baslik"]:
        sonuc["bayraklar"].append("başlık bulunamadı")
    govde = [s for s in satirlar if s not in baslik_satirlari]

    # --- Şekiller
    tablo = tablo_bul(govde)
    if tablo:
        tablo_ids = {id(s) for s in tablo["kaynak"]}
        govde = [s for s in govde if id(s) not in tablo_ids]
        sonuc["bloklar"].append(("tablo", tablo["y"], tablo["satirlar"]))
        sonuc["bayraklar"].append("tablo (sütun/satır eşleşmesini asıl slaytla karşılaştır)")
    sekil_alanlari = list(rasterlar)
    vektor_sekiller = [r for tur, r in vektor if tur == "sekil"]
    if not tablo and len(vektor) >= 5 and vektor_sekiller:
        birlesik = pymupdf.Rect(vektor_sekiller[0])
        for r in vektor_sekiller[1:]:
            birlesik |= r
        sekil_alanlari.append(birlesik)
        sonuc["bayraklar"].append("vektör şema")
    if rasterlar:
        sonuc["bayraklar"].append("görsel/şema (metni PDF'te yok, görsele bak)")
    if sekil_alanlari:
        icerde, disarda = [], []
        for s in govde:
            merkez = pymupdf.Point((s["x0"] + s["x1"]) / 2, (s["y0"] + s["y1"]) / 2)
            if s["bullet_x"] is None and any(pymupdf.Rect(a).include_rect(a) and
                                             (pymupdf.Rect(a) + (-15, -15, 15, 15)).contains(merkez)
                                             for a in sekil_alanlari):
                icerde.append(s)
            else:
                disarda.append(s)
        if icerde:
            sonuc["sekil_etiketleri"] = etiketleri_grupla(icerde)
            sonuc["bayraklar"].append("şekil içi metin (ok/ilişki bilgisi yok)")
        govde = disarda

    # --- Maddeler ve paragraflar
    govde.sort(key=lambda s: (round(s["y0"] / 2), s["x0"]))
    ogeler, mevcut = [], None
    for s in govde:
        if s["bullet_x"] is not None:
            seviye = min(range(len(bullet_seviyeleri)), key=lambda i: abs(bullet_seviyeleri[i] - s["bullet_x"])) \
                if bullet_seviyeleri else 0
            mevcut = {"tur": "madde", "seviye": seviye, "metin": s["metin"], "x": s["x0"],
                      "alt": s["y1"], "y": s["y0"]}
            ogeler.append(mevcut)
            continue
        if mevcut and abs(s["x0"] - mevcut["x"]) < 12 and s["y0"] - mevcut["alt"] < 0.9 * s["boyut"]:
            mevcut["metin"] = satir_birlestir(mevcut["metin"], s["metin"])
            mevcut["alt"] = s["y1"]
            continue
        # Madde işareti olmayan yeni paragraf: girintiye göre seviye tahmini
        seviye = 0
        onceki_maddeler = [o for o in ogeler if o["tur"] == "madde"]
        for o in reversed(onceki_maddeler):
            if abs(o["x"] - s["x0"]) < 12:
                seviye = o["seviye"]
                break
        mevcut = {"tur": "paragraf", "seviye": seviye, "metin": s["metin"], "x": s["x0"],
                  "alt": s["y1"], "y": s["y0"]}
        ogeler.append(mevcut)
    if ogeler:
        sonuc["bloklar"].append(("liste", ogeler[0]["y"], ogeler))
    sonuc["bloklar"].sort(key=lambda b: b[1])
    return sonuc


# ──────────────────────────────────────────────────────────────────────────────
# Görüntü PDF'ler (yazı katmanı olmayan: ekran görüntüsü, tarama, sayfa başına birden çok slayt)
# ──────────────────────────────────────────────────────────────────────────────

def goruntu_pdf_mi(pdf_yol: Path) -> bool:
    """Sayfaların en az %80'inde yazı katmanı yoksa ve resim varsa bu bir görüntü PDF'tir."""
    belge = pymupdf.open(pdf_yol)
    yazisiz = sum(1 for s in belge if len(s.get_text().strip()) < 20 and s.get_images())
    return len(belge) > 0 and yazisiz >= 0.8 * len(belge)


def _sayfa_goruntusu(belge, i):
    """Sayfadaki en büyük gömülü resmi özgün çözünürlüğünde alır; yoksa sayfayı 200 dpi çizer."""
    from PIL import Image
    sayfa = belge[i]
    resimler = sayfa.get_images(full=True)
    if resimler:
        xref = max(resimler, key=lambda r: r[2] * r[3])[0]
        kutu = sayfa.get_image_rects(xref)
        if kutu and kutu[0].get_area() / sayfa.rect.get_area() > 0.3:
            pix = pymupdf.Pixmap(belge, xref)
            if pix.n - pix.alpha != 3:
                pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
            if pix.alpha:
                pix = pymupdf.Pixmap(pix, 0)
            return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
    pix = sayfa.get_pixmap(dpi=200, alpha=False)
    return Image.frombytes("RGB", (pix.width, pix.height), pix.samples)


def _koyu_zemini_cevir(resim):
    """Koyu zemin üzerindeki açık yazıyı, OCR'ın iyi okuduğu koyu-yazı/açık-zemin biçimine çevirir."""
    import numpy as np
    from PIL import ImageFilter, Image
    gri = resim.convert("L")
    a = np.asarray(gri).astype(np.int16)
    zemin = np.asarray(gri.filter(ImageFilter.BoxBlur(max(8, gri.width // 40)))).astype(np.int16)
    return Image.fromarray(np.where(zemin < 110, 255 - a, a).astype(np.uint8))


def _goruntu_satirlari(resim):
    """Tüm sayfayı iki biçimde (özgün / koyu zemin çevrilmiş) okur, daha çok ve güvenilir kelime vereni seçer.
    Satırlar: x0, y0, x1, y1, h (piksel), metin, güven, kelime kutuları."""
    import pytesseract
    en_iyi = None
    for aday in (resim.convert("L"), _koyu_zemini_cevir(resim)):
        veri = pytesseract.image_to_data(aday, lang=_OCR_DURUM["diller"], config="--psm 3",
                                         output_type=pytesseract.Output.DICT)
        gruplar = {}
        for i, k in enumerate(veri["text"]):
            if not k.strip() or float(veri["conf"][i]) < 0:
                continue
            gruplar.setdefault((veri["block_num"][i], veri["par_num"][i], veri["line_num"][i]), []).append(
                (veri["left"][i], veri["top"][i], veri["left"][i] + veri["width"][i],
                 veri["top"][i] + veri["height"][i], k, float(veri["conf"][i])))
        satirlar = []
        for kel in gruplar.values():
            kel.sort()
            yuk = sorted(k[3] - k[1] for k in kel)[len(kel) // 2]
            parcalar, parca = [], [kel[0]]
            for k in kel[1:]:                       # büyük yatay boşluk = ayrı sütun
                if k[0] - parca[-1][2] > 1.8 * yuk:
                    parcalar.append(parca)
                    parca = [k]
                else:
                    parca.append(k)
            parcalar.append(parca)
            for p in parcalar:
                satirlar.append({"x0": p[0][0], "y0": min(k[1] for k in p), "x1": p[-1][2],
                                 "y1": max(k[3] for k in p), "h": yuk, "kelimeler": p,
                                 "metin": " ".join(k[4] for k in p), "guven": sum(k[5] for k in p) / len(p)})
        puan = (sum(len(s["kelimeler"]) for s in satirlar if s["guven"] >= 80),
                sum(s["guven"] for s in satirlar) / max(1, len(satirlar)))
        if en_iyi is None or puan > en_iyi[0]:
            en_iyi = (puan, satirlar)
    return sorted(en_iyi[1], key=lambda s: (s["y0"], s["x0"]))


_ISARET_KELIMELERI = {"e", "«", "¢", "¢«", "+", "*", "•", "●", "o", "°", "-", "–", "»", "©", "@", "®", "«e", "ee"}


def _madde_isareti_ayikla(satir, gri):
    """Satır başındaki madde işaretini bulur: ya OCR işareti bir kelime olarak okumuştur ("e", "¢", "•"),
    ya da yazının solunda küçük koyu bir leke vardır. İşaretin x konumunu döndürür (yoksa None)."""
    import numpy as np
    kel = satir["kelimeler"]
    if len(kel) >= 2:
        ilk, ikinci = kel[0], kel[1]
        kisa_isaret = len(ilk[4]) <= 2 and not ilk[4].isalnum() or ilk[4] in _ISARET_KELIMELERI
        ayrik = ikinci[0] - ilk[2] >= 0.5 * max(8, satir["h"])
        if kisa_isaret and (ayrik or ilk[4] in _ISARET_KELIMELERI) and (ikinci[4][:1].isupper() or ikinci[4][:1].isdigit()):
            satir["kelimeler"] = kel[1:]
            satir["metin"] = " ".join(k[4] for k in satir["kelimeler"])
            satir["x0"] = ikinci[0]
            satir["guven"] = sum(k[5] for k in satir["kelimeler"]) / len(satir["kelimeler"])
            return ilk[0]
    # Piksel kontrolü: metnin solunda yalnız duran küçük koyu leke (solunda başka yazı olmamalı)
    h = max(8, satir["h"])
    x_bas, y_orta = satir["x0"], (satir["y0"] + satir["y1"]) // 2
    x0, x1 = max(0, int(x_bas - 4.5 * h)), max(0, int(x_bas - 0.3 * h))
    y0, y1 = max(0, int(y_orta - 0.5 * h)), int(y_orta + 0.5 * h)
    if x1 - x0 < 4 or y1 <= y0:
        return None
    parca = np.asarray(gri.crop((x0, y0, x1, y1)))
    if parca.size == 0:
        return None
    sutun_koyu = (parca < 110).any(axis=0)
    if not sutun_koyu.any():
        return None
    # Koyu sütun öbekleri
    obekler, bas = [], None
    for i, v in enumerate(sutun_koyu):
        if v and bas is None:
            bas = i
        if not v and bas is not None:
            obekler.append((bas, i - 1))
            bas = None
    if bas is not None:
        obekler.append((bas, len(sutun_koyu) - 1))
    if len(obekler) != 1:
        return None                      # birden çok öbek: solunda yazı var (tablo sütunu vb.)
    a, b = obekler[0]
    if b - a + 1 > 0.8 * h or (len(sutun_koyu) - 1 - b) < 0.2 * h:
        return None
    satir_koyu = (parca[:, a:b + 1] < 110).any(axis=1)
    if satir_koyu.sum() > 0.8 * h:
        return None
    return x0 + a


def _slayt_ortasi(satirlar, genislik):
    """Slaytların yatay orta noktası (sayfa kenarındaki düğmeler vb. hariç, yazı bloklarının ortası)."""
    if not satirlar:
        return genislik / 2
    x0 = sorted(s["x0"] for s in satirlar)[len(satirlar) // 10]
    x1 = sorted(s["x1"] for s in satirlar)[-1 - len(satirlar) // 10]
    return (x0 + x1) / 2


def _temiz_satirlar(satirlar, resim):
    """Görüntüleyici düğmeleri (sayfanın sağ/sol kenarı) ve anlamsız kısa OCR parçalarını eler."""
    temiz = []
    for s in satirlar:
        harf = sum(c.isalpha() for c in s["metin"])
        if harf < 2:
            continue
        if s["x0"] > 0.92 * resim.width or s["x1"] < 0.06 * resim.width:
            continue
        if harf <= 4 and s["guven"] < 80:
            continue
        temiz.append(s)
    return temiz


def _eksik_bolgeleri_yeniden_oku(resim, satirlar, altbilgi_y):
    """Tüm sayfa okumasında Tesseract bazen bir slayt bölgesini atlar. Alt bilgilerle ayrılan bantlardan
    büyük olup çok az kelime içerenler ayrıca kırpılıp yeniden okunur."""
    sinirlar = sorted(set(int(y) for y in altbilgi_y))
    if not sinirlar:
        return satirlar
    bantlar, onceki = [], 0
    for y in sinirlar + [resim.height]:
        bantlar.append((onceki, min(resim.height, y + 8)))
        onceki = min(resim.height, y + 8)
    sonuc = list(satirlar)
    for y0, y1 in bantlar:
        if y1 - y0 < 0.2 * resim.height:
            continue
        icerde = [s for s in satirlar if y0 <= s["y0"] < y1]
        if sum(len(s["metin"].split()) for s in icerde) >= 8:
            continue
        yeni = _temiz_satirlar(_goruntu_satirlari(resim.crop((0, y0, resim.width, y1))), resim)
        for s in yeni:
            s["y0"] += y0
            s["y1"] += y0
            s["kelimeler"] = [(k[0], k[1] + y0, k[2], k[3] + y0, k[4], k[5]) for k in s["kelimeler"]]
        if sum(len(s["metin"].split()) for s in yeni) > sum(len(s["metin"].split()) for s in icerde):
            sonuc = [s for s in sonuc if not (y0 <= s["y0"] < y1)] + yeni
    return sorted(sonuc, key=lambda s: (s["y0"], s["x0"]))


def goruntu_pdf_isle(pdf_yol: Path):
    """Görüntü PDF'ten slaytları çıkarır. Her sayfa bütün olarak OCR'dan geçer; büyük yazılar başlık sayılır ve her
    başlık yeni bir slayt başlatır (sayfa başına birden çok slayt olabilir). Madde işaretleri ve seviyeler
    konumdan kurulur, sayfalar arasında tekrar eden alt bilgiler atılır."""
    belge = pymupdf.open(pdf_yol)
    sayfalar = []
    for i in range(len(belge)):
        resim = _sayfa_goruntusu(belge, i)
        sayfalar.append((resim, _temiz_satirlar(_goruntu_satirlari(resim), resim)))

    # Alt/üst bilgi: sayfaların en az %30'unda tekrar eden satırlar (sayılar yok sayılarak)
    def anahtar(t):
        return re.sub(r"[^a-z#]", "", re.sub(r"\d+", "#", t.lower()))
    sayac = Counter()
    for _, satirlar in sayfalar:
        sayac.update({anahtar(s["metin"]) for s in satirlar})
    tekrar = {k for k, n in sayac.items() if n >= max(2, 0.3 * len(sayfalar)) and len(k) >= 6}

    slaytlar, tum_isaretler = [], []
    for sayfa_no, (resim, satirlar) in enumerate(sayfalar):
        gri = _koyu_zemini_cevir(resim)
        tum_h = sorted(x["h"] for x in satirlar)
        med_h = tum_h[len(tum_h) // 2] if tum_h else 0

        def altbilgi_mi(x):
            k = anahtar(x["metin"])
            if k in tekrar:
                return True
            if x["h"] > 0.7 * med_h or len(k) < 6:
                return False
            if any(k in t or t in k for t in tekrar if len(t) >= 6):
                return True
            import difflib
            benzer = sum(1 for j, (_, diger) in enumerate(sayfalar) if j != sayfa_no and
                         any(difflib.SequenceMatcher(None, k, anahtar(d["metin"])).ratio() >= 0.75 for d in diger))
            return benzer >= max(1, 0.3 * (len(sayfalar) - 1))
        altbilgi = [x for x in satirlar if altbilgi_mi(x)]
        satirlar = _eksik_bolgeleri_yeniden_oku(resim, satirlar, [a["y1"] for a in altbilgi])
        govde = [x for x in satirlar if not altbilgi_mi(x)]
        altbilgi = [x for x in satirlar if altbilgi_mi(x)]
        if not govde:
            continue
        yukseklikler = sorted(x["h"] for x in govde)
        medyan = yukseklikler[len(yukseklikler) // 2]
        en_buyuk_h = max(x["h"] for x in govde if x["guven"] >= 80) if any(x["guven"] >= 80 for x in govde) else medyan
        genislik = resim.width
        orta = _slayt_ortasi(govde, genislik)
        govde.sort(key=lambda x: (x["y0"], x["x0"]))
        # 1) Slayt sınırları: alt bilgi satırlarının hemen sonrası ve büyük dikey boşluklar
        sinirlar = sorted({int(a["y1"]) for a in altbilgi})
        parcalar, parca = [], []
        for x in govde:
            if parca:
                bosluk = x["y0"] - max(p["y1"] for p in parca)
                sinir_gecti = any(max(p["y1"] for p in parca) <= b <= x["y0"] for b in sinirlar)
                if sinir_gecti or (not sinirlar and bosluk > 4 * medyan):
                    parcalar.append(parca)
                    parca = []
            parca.append(x)
        if parca:
            parcalar.append(parca)

        # 2) Her parçada başlık: en üst satır(lar); ortalanmış ya da gövdeden büyükse. Parça içinde ortalanmış ve
        #    belirgin büyük başka bir satır varsa (alt bilgisi olmayan düzenler) orası da yeni slayt başlatır.
        def baslik_mi(x, ilk):
            harf = sum(c.isalpha() for c in x["metin"])
            ortali = abs((x["x0"] + x["x1"]) / 2 - orta) < 0.1 * genislik
            if harf < 3 or x["guven"] < 70:
                return False
            if ilk:
                return ortali or x["h"] >= 1.15 * medyan
            return ortali and x["h"] >= 1.5 * medyan

        def yeni_slayt(baslik, x, h):
            sl = {"baslik": baslik, "satirlar": [], "sayfa": sayfa_no, "_baslik_y": x["y0"],
                  "_baslik_h": h, "resim": resim, "_son_baslik": x if baslik else None}
            slaytlar.append(sl)
            return sl

        for parca in parcalar:
            mevcut = None
            for j, x in enumerate(parca):
                if j == 0:
                    mevcut = yeni_slayt(x["metin"], x, x["h"]) if baslik_mi(x, ilk=True) else \
                        yeni_slayt("", x, medyan)
                    if mevcut["baslik"]:
                        continue
                elif mevcut["_son_baslik"] is not None and not mevcut["satirlar"] and \
                        x["y0"] - mevcut["_son_baslik"]["y1"] < 0.9 * x["h"] and \
                        abs(x["h"] - mevcut["_son_baslik"]["h"]) < 0.25 * x["h"] and baslik_mi(x, ilk=True):
                    mevcut["baslik"] = satir_birlestir(mevcut["baslik"], x["metin"])   # çok satırlı başlık
                    mevcut["_son_baslik"] = x
                    continue
                elif baslik_mi(x, ilk=False) and x["h"] >= 0.75 * en_buyuk_h and \
                        (not sinirlar or x["y0"] - parca[j - 1]["y1"] >= 2.0 * x["h"]):
                    # Alt bilgi okunamamış olsa da büyük boşluktan sonra gelen ortalanmış büyük satır yeni slayttır
                    mevcut = yeni_slayt(x["metin"], x, x["h"])
                    continue
                x["bullet_x"] = _madde_isareti_ayikla(x, gri)
                if x["guven"] < 75:                 # şema etiketleri vb.: güvenilmez OCR satırı
                    mevcut["_atilan"] = mevcut.get("_atilan", 0) + 1
                    continue
                mevcut["satirlar"].append(x)
        # Her alt bilgi bir slaytı kapatır: üstünde hiç slayt açılmamış bir alt bilgi (yazısı okunamayan şema
        # slaytı) varsa o bant başlıksız bir görsel slayt olarak eklenir — hiçbir slayt kaybolmaz.
        altbilgi_y = sorted(kume([a["y0"] for a in altbilgi], 20)) if altbilgi else []
        onceki_sinir = 0
        for ay, _ in altbilgi_y:
            bantta = [sl for sl in slaytlar if sl["sayfa"] == sayfa_no and onceki_sinir <= sl["_baslik_y"] < ay]
            if not bantta and ay - onceki_sinir > 0.15 * resim.height:
                bos = {"baslik": "", "satirlar": [], "sayfa": sayfa_no, "_baslik_y": onceki_sinir + 10,
                       "_baslik_h": medyan, "resim": resim, "_son_baslik": None, "_yazisiz": True}
                sira = next((k for k, sl in enumerate(slaytlar) if sl["sayfa"] == sayfa_no and sl["_baslik_y"] > ay),
                            None)
                if sira is None:
                    sira = max([k + 1 for k, sl in enumerate(slaytlar) if sl["sayfa"] <= sayfa_no], default=0)
                slaytlar.insert(sira, bos)
            onceki_sinir = ay + 10
        # Slaytların alt sınırı: kendi alt bilgisi ya da bir sonraki slaytın başı
        sayfa_slaytlari = [sl for sl in slaytlar if sl["sayfa"] == sayfa_no]
        for k, sl in enumerate(sayfa_slaytlari):
            sonraki = sayfa_slaytlari[k + 1]["_baslik_y"] if k + 1 < len(sayfa_slaytlari) else resim.height
            alt = [a for a in altbilgi if sl["_baslik_y"] < a["y0"] < sonraki]
            sl["_alt_var"] = bool(alt)
            sl["_alt_y"] = (max(a["y1"] for a in alt) + int(0.6 * sl["_baslik_h"])) if alt else \
                (sonraki - int(0.6 * sl["_baslik_h"]) if k + 1 < len(sayfa_slaytlari) else resim.height)
            no = next((re.search(r"(?i)slide\s*(\d+)", a["metin"]) for a in alt
                       if re.search(r"(?i)slide\s*(\d+)", a["metin"])), None)
            sl["_slayt_no"] = int(no.group(1)) if no else None

    # Aynı sayfada başlıksız ve çok az yazılı parça: önceki slaytın (şema etiketi gibi) kopmuş bir parçasıdır
    birlesik = []
    for sl in slaytlar:
        kelime = sum(len(x["metin"].split()) for x in sl["satirlar"])
        if not sl["baslik"] and kelime < 6 and birlesik and birlesik[-1]["sayfa"] == sl["sayfa"] \
                and not birlesik[-1].get("_alt_var"):      # önceki slayt kendi alt bilgisiyle kapanmışsa bu ayrı slayttır
            birlesik[-1]["satirlar"] += sl["satirlar"]
            birlesik[-1]["_atilan"] = birlesik[-1].get("_atilan", 0) + sl.get("_atilan", 0)
            birlesik[-1]["_alt_y"] = max(birlesik[-1].get("_alt_y", 0), sl.get("_alt_y", 0))
            continue
        birlesik.append(sl)
    slaytlar = birlesik
    sonuc = []
    for sl in slaytlar:
        isaretler = [s["bullet_x"] for s in sl["satirlar"] if s["bullet_x"] is not None]
        tolerans = max(8, 0.8 * (sorted(s["h"] for s in sl["satirlar"])[len(sl["satirlar"]) // 2] if sl["satirlar"] else 10))
        seviyeler = [m for m, n in kume(isaretler, tolerans)] if isaretler else []
        ogeler, son = [], None
        for s in sl["satirlar"]:
            if s["bullet_x"] is not None:
                seviye = min(range(len(seviyeler)), key=lambda i: abs(seviyeler[i] - s["bullet_x"])) if seviyeler else 0
                son = {"tur": "madde", "seviye": seviye, "metin": s["metin"], "x": s["x0"], "alt": s["y1"],
                       "y": s["y0"], "h": s["h"]}
                ogeler.append(son)
            elif son and abs(s["x0"] - son["x"]) < 0.9 * s["h"] and s["y0"] - son["alt"] < 1.2 * s["h"]:
                son["metin"] = satir_birlestir(son["metin"], s["metin"])
                son["alt"] = s["y1"]
            else:
                seviye = next((o["seviye"] for o in reversed(ogeler) if abs(o["x"] - s["x0"]) < 0.9 * s["h"]), 0)
                son = {"tur": "paragraf", "seviye": seviye, "metin": s["metin"], "x": s["x0"], "alt": s["y1"],
                       "y": s["y0"], "h": s["h"]}
                ogeler.append(son)
        madde_sayisi = sum(1 for o in ogeler if o["tur"] == "madde")
        sol_kenarlar = kume([s["x0"] for s in sl["satirlar"]], 0.8 * max(10, sorted(s["h"] for s in sl["satirlar"])[len(sl["satirlar"]) // 2])) if sl["satirlar"] else []
        ort_kelime = (sum(len(s["metin"].split()) for s in sl["satirlar"]) / len(sl["satirlar"])) if sl["satirlar"] else 0
        sema_duzeni = madde_sayisi == 0 and sl["satirlar"] and (len(sol_kenarlar) >= 3 or ort_kelime < 4)
        guvenler = [s["guven"] for s in sl["satirlar"]]
        kelime = sum(len(s["metin"].split()) for s in sl["satirlar"])
        yapi = {"baslik": sl["baslik"], "bloklar": [("liste", 0, ogeler)] if ogeler else [], "sekil_etiketleri": [],
                "bayraklar": [], "not": "", "duzeltmeler": [], "gorsel_alanlari": [], "kaynak_metin": None,
                "ocr_metni": True}
        # Tablo: sütunlu satırlar varsa tablo olarak kur
        adaylar = [dict(s, boyut=s["h"] * 1.35, _kaynak=id(s)) for s in sl["satirlar"]]
        for a in adaylar:
            a["bullet_x"] = None
        tablo = tablo_bul(adaylar)
        if tablo:
            bolunmus = []
            for a in adaylar:
                kel = a.get("kelimeler") or []
                kesim = next((c for c in tablo["sutunlar"][1:] if a["x0"] < c - 15 and a["x1"] > c + 30), None)
                if kesim is None or not kel:
                    bolunmus.append(a)
                    continue
                sol = [k for k in kel if k[0] < kesim - 10 and any(ch.isalnum() for ch in k[4])]
                while sol and len(sol[-1][4]) == 1 and sol[-1][4] not in ("a", "A", "I"):
                    sol.pop()              # sütun sınırındaki tek karakterlik OCR kalıntısı ("o")
                sag = [k for k in kel if k[0] >= kesim - 10]
                for grup in (sol, sag):
                    if grup:
                        bolunmus.append(dict(a, x0=grup[0][0], x1=grup[-1][2], kelimeler=grup,
                                             metin=" ".join(k[4] for k in grup)))
            if len(bolunmus) != len(adaylar):
                adaylar = bolunmus
                tablo = tablo_bul(adaylar) or tablo
        if tablo and len(tablo["satirlar"]) >= 3 and len(tablo["kaynak"]) >= 0.6 * len(adaylar):
            kullanilan = {a["_kaynak"] for a in tablo["kaynak"]}
            yapi["bloklar"] = [("tablo", 0, tablo["satirlar"])]
        if sema_duzeni and yapi["bloklar"] and yapi["bloklar"][0][0] == "liste":
            yapi["sekil_etiketleri"] = [o["metin"] for o in ogeler]
            yapi["bloklar"] = []
            yapi["bayraklar"].append("görsel/şema (OCR yalnızca etiketleri okur; ilişkiler için görsele bak)")
        # Kırpım: başlığın biraz üstünden slaytın alt sınırına
        ust = max(0, int(sl["_baslik_y"] - 2.2 * sl["_baslik_h"]))
        yapi["kirpim"] = sl["resim"].crop((0, ust, sl["resim"].width, min(sl["resim"].height, sl["_alt_y"])))
        yapi["kaynak_konum"] = f"PDF s.{sl['sayfa'] + 1}" + (f", slayt {sl['_slayt_no']}" if sl["_slayt_no"] else "")
        dusuk = (min(guvenler) if guvenler else 0) < 75 or (sum(guvenler) / len(guvenler) if guvenler else 0) < 90
        if sl.get("_yazisiz"):
            yapi["bayraklar"].append("görsel: bu slaytta okunabilir yazı bulunamadı (şema/resim), görsele bak")
        elif kelime < 15 or sl.get("_atilan", 0) >= 2:
            yapi["bayraklar"].append("görsel/şema olabilir (az ya da güvenilmez yazı; görsele bak)")
        if dusuk:
            yapi["bayraklar"].append("görsel: bazı satırların OCR güveni düşük, görsele bak")
        if not sl["baslik"]:
            yapi["bayraklar"].append("başlık bulunamadı")
        yapi["bayraklar"].append("metin görüntüden OCR ile okundu")
        sonuc.append(yapi)
    return sonuc


def tekrar_eden_satirlar(belge):
    """Sayfaların en az yarısında tekrar eden üst/alt bilgi satırları (sayılar yok sayılarak)."""
    sayac = Counter()
    for sayfa in belge:
        H = sayfa.rect.height
        gorulen = set()
        for blok in sayfa.get_text("dict")["blocks"]:
            if blok["type"] != 0:
                continue
            for satir in blok["lines"]:
                t = "".join(s["text"] for s in satir["spans"]).strip()
                y = satir["bbox"][1]
                if t and (y > 0.88 * H or y < 0.08 * H):
                    gorulen.add(normal_metin(t))
        sayac.update(gorulen)
    esik = max(2, len(belge) // 2)
    return {t for t, n in sayac.items() if n >= esik}


def pdf_isle(pdf_yol: Path):
    belge = pymupdf.open(pdf_yol)
    tekrar = tekrar_eden_satirlar(belge)

    def ustbilgi_mi(s, H):
        return (s["y0"] > 0.88 * H or s["y0"] < 0.08 * H) and (
            normal_metin(s["metin"]) in tekrar or
            any(normal_metin(s["metin"]) in t or t in normal_metin(s["metin"]) for t in tekrar if len(t) > 8))

    ham = []
    tum_bulletlar = []
    for sayfa in belge:
        duzeltmeler = []
        satirlar, isaretler = sayfa_satirlari(sayfa, duzeltmeler)
        H = sayfa.rect.height
        ust_alt = [s for s in satirlar if ustbilgi_mi(s, H)]
        satirlar = [s for s in satirlar if s not in ust_alt]
        isaretler = [i for i in isaretler if not (i["y0"] > 0.88 * H)]
        sahipsiz = isaretleri_eslestir(satirlar, isaretler)
        tum_bulletlar += [s["bullet_x"] for s in satirlar if s["bullet_x"] is not None]
        ham.append((sayfa, satirlar, sahipsiz, duzeltmeler, ust_alt))

    sozluk_birlestir([s for _, satirlar, _, d, _ in ham for s in satirlar],
                     {id(s): d for _, satirlar, _, d, _ in ham for s in satirlar}, EK_SOZLUK)

    # Belge geneli madde seviyeleri (sık görülen girinti konumları)
    kumeler = kume(tum_bulletlar, 10)
    esik = max(2, int(0.01 * len(tum_bulletlar)))
    seviyeler = [m for m, n in kumeler if n >= esik] or [m for m, _ in kumeler]

    slaytlar = []
    for sayfa, satirlar, sahipsiz, duzeltmeler, ust_alt in ham:
        yapi = pdf_sayfasi_yapilandir(sayfa, satirlar, seviyeler)
        if sahipsiz:
            yapi["bayraklar"].append(f"{sahipsiz} madde işareti bir satıra bağlanamadı")
        if duzeltmeler:
            yapi["bayraklar"].append(f"{len(duzeltmeler)} hatalı boşluk düzeltildi")
        if not satirlar:
            yapi["bayraklar"].append("hiç metin yok (taranmış olabilir; görsele bak)")
        # Bağımsız karşılaştırma metni: PyMuPDF'in düz metin çıkarımı, üst/alt bilgi hariç
        H = sayfa.rect.height
        duz = []
        for blok in sayfa.get_text("dict")["blocks"]:
            if blok["type"] != 0:
                continue
            for satir in blok["lines"]:
                t = "".join(s["text"] for s in satir["spans"])
                sahte = {"y0": satir["bbox"][1], "metin": t}
                if not ustbilgi_mi(sahte, H):
                    duz.append(t)
        yapi["kaynak_metin"] = "\n".join(duz)
        yapi["duzeltmeler"] = duzeltmeler
        yapi["not"] = ""
        slaytlar.append(yapi)
    ornekler = {}
    for _, _, _, _, ua in ham:
        for s in ua:
            ornekler.setdefault(normal_metin(s["metin"]), s["metin"])
    silinen = list(ornekler.values())[:5]
    return slaytlar, silinen


# ──────────────────────────────────────────────────────────────────────────────
# PPTX'ten yapı çıkarma
# ──────────────────────────────────────────────────────────────────────────────

def pptx_isle(pptx_yol: Path):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

    sunu = Presentation(str(pptx_yol))
    slayt_gen, slayt_yuk = (sunu.slide_width or 9144000), (sunu.slide_height or 6858000)
    slayt_alani = slayt_gen * slayt_yuk
    slaytlar = []

    def metin_temizle(t):
        return re.sub(r"\s+", " ", (t or "").replace("\x0b", " ")).strip()

    for slayt in sunu.slides:
        gizli = slayt._element.get("show") == "0"
        yapi = {"baslik": "", "bloklar": [], "sekil_etiketleri": [], "bayraklar": [], "not": "",
                "duzeltmeler": [], "gizli": gizli, "gorsel_alanlari": []}

        def alan_ekle(sekil, grup_kutusu):
            sol, ust, gen, yuk = grup_kutusu or (sekil.left or 0, sekil.top or 0, sekil.width or 0, sekil.height or 0)
            kutu = (sol / slayt_gen, ust / slayt_yuk, (sol + gen) / slayt_gen, (ust + yuk) / slayt_yuk)
            alan_ekle_tekil(yapi["gorsel_alanlari"], kutu)
        baglayici_var = False
        sekil_metinleri = []
        gorsel_var = False

        def gez(sekiller, grup_icinde=False, grup_kutusu=None):
            nonlocal baglayici_var, gorsel_var
            for sekil in sorted(sekiller, key=lambda s: ((s.top or 0), (s.left or 0))):
                tur = sekil.shape_type
                y = sekil.top or 0
                if tur == MSO_SHAPE_TYPE.GROUP:
                    # Grup içindeki şekillerin koordinatları grubun kendi düzlemindedir; slayttaki yer için
                    # en dıştaki grubun kutusu kullanılır.
                    gez(sekil.shapes, True, grup_kutusu or (sekil.left or 0, sekil.top or 0,
                                                            sekil.width or 0, sekil.height or 0))
                    continue
                if getattr(sekil, "is_placeholder", False) and sekil.placeholder_format.type in (
                        PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE):
                    yapi["baslik"] = metin_temizle(sekil.text_frame.text) if sekil.has_text_frame else ""
                    continue
                if getattr(sekil, "is_placeholder", False) and sekil.placeholder_format.type in (
                        PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE):
                    continue
                if tur in (MSO_SHAPE_TYPE.LINE,) or sekil.__class__.__name__ == "Connector":
                    baglayici_var = True
                    continue
                if tur == MSO_SHAPE_TYPE.PICTURE or (getattr(sekil, "is_placeholder", False) and
                                                    sekil.__class__.__name__ == "PlaceholderPicture"):
                    alan = (sekil.width or 0) * (sekil.height or 0)
                    if alan >= 0.03 * slayt_alani:     # logo/simge gibi küçük resimler sayılmaz
                        gorsel_var = True
                        alan_ekle(sekil, grup_kutusu)
                    continue
                if getattr(sekil, "has_table", False) and sekil.has_table:
                    satirlar = [[metin_temizle(h.text) for h in r.cells] for r in sekil.table.rows]
                    yapi["bloklar"].append(("tablo", y, satirlar))
                    continue
                if getattr(sekil, "has_chart", False) and sekil.has_chart:
                    alan_ekle(sekil, grup_kutusu)
                    try:
                        grafik = sekil.chart
                        kategoriler = [str(k) for k in grafik.plots[0].categories]
                        satirlar = [["Seri"] + kategoriler]
                        for seri in grafik.series:
                            satirlar.append([seri.name] + [("" if v is None else f"{v:g}") for v in seri.values])
                        yapi["bloklar"].append(("tablo", y, satirlar))
                        yapi["bayraklar"].append("grafik (değerler sunudan okundu)")
                    except Exception:
                        yapi["bayraklar"].append("grafik (değerler okunamadı, görsele bak)")
                    continue
                if not getattr(sekil, "has_text_frame", False) or not sekil.has_text_frame:
                    if tur in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
                               MSO_SHAPE_TYPE.FREEFORM, MSO_SHAPE_TYPE.AUTO_SHAPE):
                        if tur in (MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT):
                            gorsel_var = True
                            alan_ekle(sekil, grup_kutusu)
                    continue
                paragraflar = [(p.level, metin_temizle("".join(r.text for r in p.runs) or p.text if hasattr(p, "text") else ""))
                               for p in sekil.text_frame.paragraphs]
                paragraflar = [(l, t) for l, t in paragraflar if t]
                if not paragraflar:
                    continue
                govde_yer_tutucu = getattr(sekil, "is_placeholder", False)
                duz_metin_kutusu = tur == MSO_SHAPE_TYPE.TEXT_BOX
                if grup_icinde or (not govde_yer_tutucu and not duz_metin_kutusu):
                    # Otomatik şekil (kutu, daire...) içindeki metin: şema etiketi olabilir, sonra karar verilir
                    sekil_metinleri.append((y, paragraflar, grup_icinde))
                    continue
                ogeler = [{"tur": "madde" if govde_yer_tutucu else "paragraf", "seviye": l, "metin": t, "y": y}
                          for l, t in paragraflar]
                yapi["bloklar"].append(("liste", y, ogeler))

        gez(slayt.shapes)
        # Az sayıda otomatik şekil + bağlayıcı yoksa bunlar büyük ihtimalle düz metin kutusudur
        # Bağlayıcı/grup yoksa ve az sayıda şekil varsa bunlar metin kutusudur (ör. .ppt dönüşümünde
        # yer tutucusunu kaybetmiş gövde metni): paragraf ve seviyeler korunur.
        sema = baglayici_var or any(g for _, _, g in sekil_metinleri) or len(sekil_metinleri) >= 3
        if sekil_metinleri and not sema:
            for y, paragraflar, _ in sekil_metinleri:
                madde = len(paragraflar) > 1 or any(l > 0 for l, _ in paragraflar)
                yapi["bloklar"].append(("liste", y, [{"tur": "madde" if madde else "paragraf", "seviye": l,
                                                      "metin": t, "y": y} for l, t in paragraflar]))
        elif sekil_metinleri:
            yapi["sekil_etiketleri"] = [" ".join(t for _, t in p) for _, p, _ in
                                        sorted(sekil_metinleri, key=lambda x: x[0])]
            yapi["bayraklar"].append("şema/şekil içi metin (ok/ilişki bilgisi yok, görsele bak)")
        if gorsel_var:
            yapi["bayraklar"].append("görsel (içindeki yazı çıkarılamaz, görsele bak)")
        yapi["bloklar"].sort(key=lambda b: b[1])
        if slayt.has_notes_slide:
            yapi["not"] = (slayt.notes_slide.notes_text_frame.text or "").strip() \
                if slayt.notes_slide.notes_text_frame else ""
        if not yapi["baslik"]:
            yapi["bayraklar"].append("başlık yer tutucusu yok")
        slaytlar.append(yapi)
    return slaytlar


# ──────────────────────────────────────────────────────────────────────────────
# Markdown ve rapor üretimi
# ──────────────────────────────────────────────────────────────────────────────

BOS_ESIGI = 0.995   # görsel alanının bu oranı tek renkse "boş çizildi" sayılır
BOS_UYARISI = "görsel BOŞ çizildi (dönüştürücü bu nesneyi çizemedi)"


def baskin_renk_orani(pix, kutu):
    """Görüntüde verilen alanın (0-1 koordinat) en yaygın renginin oranı.
    Boş çizilmiş bir nesne tek renkli bir kutu olarak görünür (oran ≈ 1.0);
    gerçek şema ve grafiklerde bu oran belirgin biçimde düşüktür."""
    W, H, n, veri = pix.width, pix.height, pix.n, pix.samples
    x0, y0, x1, y1 = kutu
    dx, dy = (x1 - x0) * 0.04, (y1 - y0) * 0.04          # kenar çizgilerini sayma
    X0, Y0 = max(0, int((x0 + dx) * W)), max(0, int((y0 + dy) * H))
    X1, Y1 = min(W, int((x1 - dx) * W)), min(H, int((y1 - dy) * H))
    if X1 - X0 < 4 or Y1 - Y0 < 4 or n < 3:
        return None
    sayac = Counter()
    for y in range(Y0, Y1, 2):
        satir = y * pix.stride
        for x in range(X0, X1, 2):
            i = satir + x * n
            sayac[(veri[i] >> 3, veri[i + 1] >> 3, veri[i + 2] >> 3)] += 1
    toplam = sum(sayac.values())
    return sayac.most_common(1)[0][1] / toplam if toplam else None


def _ortusme(a, b):
    """İki kutunun kesişim/birleşim oranı (0-1)."""
    ix = max(0, min(a[2], b[2]) - max(a[0], b[0]))
    iy = max(0, min(a[3], b[3]) - max(a[1], b[1]))
    kes = ix * iy
    bir = (a[2] - a[0]) * (a[3] - a[1]) + (b[2] - b[0]) * (b[3] - b[1]) - kes
    return kes / bir if bir > 0 else 0


def alan_ekle_tekil(alanlar, kutu):
    """Aynı resmi iki kaynak (sunu yapısı + PDF) ayrı ayrı bulabilir; neredeyse çakışanlar tek sayılır."""
    if not any(_ortusme(kutu, k) > 0.7 for k in alanlar):
        alanlar.append(kutu)


def alanlari_olc(belge, sayfa_no, alanlar):
    """Sayfadaki görsel alanlarının doluluk oranları (1 - baskın renk oranı)."""
    kucuk = belge[sayfa_no].get_pixmap(dpi=72, alpha=False)
    sonuc = []
    for kutu in alanlar:
        oran = baskin_renk_orani(kucuk, kutu)
        if oran is not None:
            sonuc.append(round(1 - oran, 4))
    return sonuc


def bos_mu(doluluk):
    return any(d <= 1 - BOS_ESIGI for d in doluluk)


class YedekCizimler:
    """İlk çizim boş çıkarsa denenecek diğer çiziciler. Yalnızca gerekince (tembel) üretilir,
    üretilen PDF'ler desteler boyunca bir kez açılıp saklanır."""

    def __init__(self, adaylar):
        self.adaylar = list(adaylar)      # [(araç adı, pdf üreten fonksiyon)]
        self.hazir = []                   # [(araç adı, açık belge)]

    def __iter__(self):
        i = 0
        while True:
            if i < len(self.hazir):
                yield self.hazir[i]
                i += 1
                continue
            if not self.adaylar:
                return
            arac, uret = self.adaylar.pop(0)
            try:
                self.hazir.append((arac, pymupdf.open(uret())))
            except Exception as e:
                print(f"   ⚠ Yedek çizici {arac} kullanılamadı ({e})")


_OCR_DURUM = {}


def ocr_hazir():
    """Tesseract kullanılabilir mi? (bir kez kontrol edilir)"""
    if "hazir" not in _OCR_DURUM:
        try:
            import pytesseract
            diller = set(pytesseract.get_languages(config=""))
            secili = "+".join(d for d in OCR_DILLERI.split("+") if d in diller)
            _OCR_DURUM.update(hazir=bool(secili), diller=secili)
        except Exception:
            _OCR_DURUM.update(hazir=False, diller="")
    return OCR and _OCR_DURUM["hazir"]


def ocr_oku(resim, pt_basina_px):
    """Resimdeki yazıyı satırlar hâlinde okur. Satırlar, PDF satırlarıyla aynı biçimde (nokta cinsinden
    koordinat) döner; böylece tablo algılayıcı aynen kullanılabilir. Büyük yatay boşluklar ayrı sütun sayılır."""
    import pytesseract
    veri = pytesseract.image_to_data(resim.convert("RGB"), lang=_OCR_DURUM["diller"], config="--psm 3",
                                     output_type=pytesseract.Output.DICT)
    gruplar = {}
    for i, kelime in enumerate(veri["text"]):
        guven = float(veri["conf"][i])
        if not kelime.strip() or guven < 0:
            continue
        anahtar = (veri["block_num"][i], veri["par_num"][i], veri["line_num"][i])
        x0, y0 = veri["left"][i] / pt_basina_px, veri["top"][i] / pt_basina_px
        gruplar.setdefault(anahtar, []).append((x0, y0, x0 + veri["width"][i] / pt_basina_px,
                                                y0 + veri["height"][i] / pt_basina_px, kelime, guven))
    parcalar = []
    for kel in gruplar.values():
        kel.sort()
        h = sorted(k[3] - k[1] for k in kel)[len(kel) // 2]
        parca = [kel[0]]
        for k in kel[1:]:
            if k[0] - parca[-1][2] > 1.6 * h:
                parcalar.append(parca)
                parca = [k]
            else:
                parca.append(k)
        parcalar.append(parca)
    satirlar, guvenler = [], []
    for p in parcalar:
        guvenler += [k[5] for k in p]
        satirlar.append({"x0": p[0][0], "y0": min(k[1] for k in p), "x1": p[-1][2], "y1": max(k[3] for k in p),
                         "boyut": sorted(k[3] - k[1] for k in p)[len(p) // 2] * 1.35,
                         "metin": " ".join(k[4] for k in p), "bullet_x": None})
    return satirlar, (sum(guvenler) / len(guvenler) if guvenler else 0.0)


def _cumle_mi(metin):
    """Güvenilir sayılacak kadar uzun, gerçek kelimelerden oluşan bir satır mı? (en az 4 kelime)"""
    return sum(1 for k in metin.split() if sum(c.isalpha() for c in k) >= 2) >= 4


def _temizle(metin):
    return " ".join(k for k in metin.split() if any(c.isalpha() or c.isdigit() for c in k))


def ocr_yapilandir(satirlar):
    """OCR sonucundan yalnızca güvenilir kısımları alır:
    - tablo olarak kurulabilen bölüm (terim/tanım tabloları gibi),
    - en az 4 gerçek kelimelik cümle satırları.
    Kısa şema etiketleri alınmaz: Tesseract şemalarda etiket kaçırıp çöp üretebildiği için eksik bir
    etiket listesi yanıltıcı olur; şemanın kesin kaynağı görseldir."""
    for s in satirlar:
        s["metin"] = _temizle(s["metin"])
    satirlar = [s for s in satirlar if s["metin"]]
    tablo = tablo_bul(satirlar)
    if tablo:   # şema etiketlerinden sahte tablo çıkmasın: en az 3 satır ve hücrelerin çoğu dolu
        hucreler = [h for r in tablo["satirlar"] for h in r]
        if len(tablo["satirlar"]) < 3 or sum(1 for h in hucreler if h) < 0.75 * len(hucreler):
            tablo = None
    kalan = satirlar
    if tablo:
        idler = {id(s) for s in tablo["kaynak"]}
        kalan = [s for s in satirlar if id(s) not in idler]
    cumleler = [e for e in etiketleri_grupla(kalan) if _cumle_mi(e)] if kalan else []
    return {"tablo": tablo["satirlar"] if tablo else None, "etiketler": cumleler}


def ocr_kelime_onar(metin, sozluk):
    """OCR'ın ya da eski çizimin böldüğü kelimeleri ("Defini tion", "probabilit y") sözlükle birleştirir.
    Birleştirme koşulu: bitişik hâl sözlükte var VE (parçalardan biri hiç geçmiyor YA DA bitişik hâl,
    nadir olan parçadan en az 10 kat daha sık). "in to", "a gain" gibi iki gerçek kelime birleştirilmez."""
    if not sozluk:
        return metin
    desen = re.compile(r"\b([A-Za-zÇĞİÖŞÜçğıöşü]+) (?=([a-zçğıöşü]+)\b)")   # örtüşen çiftler için ileri bakış
    degisti = True
    while degisti:
        degisti = False
        for m in desen.finditer(metin):
            sol, sag = m.group(1), m.group(2)
            birlesik, a, b = (sol + sag).lower(), sozluk.get(sol.lower(), 0), sozluk.get(sag, 0)
            if sozluk.get(birlesik, 0) > 0 and (a == 0 or b == 0 or sozluk[birlesik] >= 10 * min(a, b)):
                metin = metin[:m.start()] + sol + metin[m.end():]
                degisti = True
                break
    return metin


def ocr_metin_duzelt(metin, sozluk, esik=3):
    """Düşük çözünürlüklü OCR'da sık görülen iki hatayı sözlükle düzeltir (yalnızca sözlük büyükse):
    - bitişen kelimeler: "thecosts" → "the costs", "i tis" → "it is"
    - tek harf hatası: "conseguent" → "consequent" (tek ve yaygın aday varsa, 6+ harfli kelimede)
    Sözlükte olan (bilinen) kelimelere asla dokunulmaz."""
    if not sozluk or len(sozluk) < 1000:
        return metin

    def bilinen(k):
        return sozluk.get(k.lower(), 0) >= esik

    def ayir(k):
        en_iyi = None
        for i in range(1, len(k)):
            a, b = k[:i], k[i:]
            if (len(a) > 1 or a.lower() in ("a", "i")) and (len(b) > 1 or b.lower() in ("a", "i")) \
                    and bilinen(a) and bilinen(b):
                puan = min(sozluk[a.lower()], sozluk[b.lower()])
                if en_iyi is None or puan > en_iyi[0]:
                    en_iyi = (puan, a, b)
        return en_iyi

    harfler = "abcdefghijklmnopqrstuvwxyz"

    def tek_harf(k):
        if len(k) < 6:
            return None
        adaylar = set()
        kl = k.lower()
        for i in range(len(kl) + 1):
            for h in harfler:
                adaylar.add(kl[:i] + h + kl[i:])
            if i < len(kl):
                adaylar.add(kl[:i] + kl[i + 1:])
                for h in harfler:
                    adaylar.add(kl[:i] + h + kl[i + 1:])
        uygun = [a for a in adaylar if sozluk.get(a, 0) >= esik]
        if len(uygun) == 1:
            a = uygun[0]
            return a.capitalize() if k[:1].isupper() else a
        return None

    parcalar = re.findall(r"[A-Za-zÇĞİÖŞÜçğıöşü]+|[^A-Za-zÇĞİÖŞÜçğıöşü]+", metin)
    for j, p in enumerate(parcalar):            # "İtis" (Türkçe İ) → "Itis": İngilizce metinde dil paketi karışması
        if "İ" in p and not bilinen(p.replace("İ", "i")) and (bilinen(p.replace("İ", "I")) or ayir(p.replace("İ", "I"))):
            parcalar[j] = p.replace("İ", "I")
    i = 0
    while i < len(parcalar):
        p = parcalar[i]
        if p[:1].isalpha() and not bilinen(p):
            # önce bir sonraki kelimeyle birleştirip yeniden bölmeyi dene ("i tis" → "it is")
            if i + 2 < len(parcalar) and parcalar[i + 1] == " " and parcalar[i + 2][:1].isalpha():
                sonuc = ayir(p + parcalar[i + 2])
                if sonuc and not bilinen(parcalar[i + 2]) or (sonuc and len(p) == 1):
                    parcalar[i:i + 3] = [sonuc[1], " ", sonuc[2]]
                    i += 3
                    continue
            if i >= 2 and parcalar[i - 1] == " " and parcalar[i - 2][:1].isalpha() and len(parcalar[i - 2]) <= 2:
                sonuc = ayir(parcalar[i - 2] + p)
                if sonuc:
                    parcalar[i - 2:i + 1] = [sonuc[1], " ", sonuc[2]]
                    i += 1
                    continue
            sonuc = ayir(p)
            if sonuc:
                parcalar[i] = sonuc[1] + " " + sonuc[2]
            else:
                duz = tek_harf(p)
                if duz:
                    parcalar[i] = duz
        i += 1
    return "".join(parcalar)


def _ocr_onar(sonuc, sozluk):
    if sonuc["tablo"]:
        sonuc["tablo"] = [[ocr_kelime_onar(h, sozluk) for h in r] for r in sonuc["tablo"]]
    sonuc["etiketler"] = [ocr_kelime_onar(e, sozluk) for e in sonuc["etiketler"]]
    return sonuc


def _en_iyi_okuma(adaylar):
    """Farklı ölçeklerde yapılan okumalar arasından seçim: tablo kurulabilen > daha çok kelime > daha yüksek güven.
    (Gerçek slaytlarda Tesseract'ın bazı satır başı kelimelerini ölçeğe bağlı olarak düşürdüğü ölçüldü.)"""
    en_iyi = None
    for resim, olcek in adaylar:
        satirlar, guven = ocr_oku(resim, olcek)
        if not satirlar:
            continue
        ham_kelime = sum(len(s["metin"].split()) for s in satirlar)
        sonuc = ocr_yapilandir(satirlar)
        sonuc["guven"] = guven
        sonuc["kelime"] = sum(len(h.split()) for r in (sonuc["tablo"] or []) for h in r) + \
            sum(len(e.split()) for e in sonuc["etiketler"])
        puan = (sonuc["tablo"] is not None, ham_kelime, guven)
        if en_iyi is None or puan > en_iyi[0]:
            en_iyi = (puan, sonuc)
    return en_iyi[1] if en_iyi and en_iyi[1]["kelime"] else None


def slayt_ocr(belge, sayfa_no, alanlar, ek_resim_yollari, sozluk=None):
    """Görsel alanlarındaki (ya da dosyadan çıkarılan resimlerdeki) yazıyı iki ölçekte okur, en iyisini alır."""
    from PIL import Image
    import io
    sonuclar = []
    if ek_resim_yollari:
        for yol in ek_resim_yollari:
            resim = Image.open(yol).convert("RGB")
            olcek = resim.width / 720          # resmin ~720 nokta genişliğinde olduğu varsayılır
            adaylar = [(resim, olcek), (resim.resize((int(resim.width * 1.5), int(resim.height * 1.5)),
                                                     Image.LANCZOS), olcek * 1.5)]
            sonuc = _en_iyi_okuma(adaylar)
            if sonuc:
                sonuclar.append(_ocr_onar(sonuc, sozluk))
        return sonuclar
    sayfa = belge[sayfa_no]
    R = sayfa.rect
    for x0, y0, x1, y1 in alanlar:
        kirp = pymupdf.Rect(x0 * R.width, y0 * R.height, x1 * R.width, y1 * R.height) & R
        if kirp.is_empty or kirp.width < 20 or kirp.height < 20:
            continue
        adaylar = []
        for dpi in (150, 220):
            pix = sayfa.get_pixmap(dpi=dpi, clip=kirp, alpha=False)
            adaylar.append((Image.open(io.BytesIO(pix.tobytes("png"))), dpi / 72))
        sonuc = _en_iyi_okuma(adaylar)
        if sonuc:
            sonuclar.append(_ocr_onar(sonuc, sozluk))
    return sonuclar


BITMAP_UZANTILARI = {"png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"}


def resim_kaydet(veri: bytes, uzanti: str, hedef: Path, gecici: Path):
    """Gömülü resmi PNG olarak kaydeder; boş/çok küçükse ya da çevrilemezse False döner."""
    uzanti = uzanti.lower().lstrip(".")
    try:
        if uzanti in BITMAP_UZANTILARI:
            from PIL import Image
            import io
            resim = Image.open(io.BytesIO(veri))
            if min(resim.size) < 48:
                return False
            resim.convert("RGB").save(hedef)
        elif uzanti in ("emf", "wmf"):      # vektör önizlemeler: LibreOffice ile PNG'ye çevrilir
            gecici.mkdir(parents=True, exist_ok=True)
            kaynak = gecici / f"{hedef.stem}.{uzanti}"
            kaynak.write_bytes(veri)
            shutil.move(str(soffice_cevir(kaynak, "png", gecici / "png")), hedef)
        else:
            return False
        pix = pymupdf.Pixmap(str(hedef))
        if pix.alpha or pix.n < 3:
            pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
        oran = baskin_renk_orani(pix, (0, 0, 1, 1))
        if oran is None or oran >= BOS_ESIGI:
            hedef.unlink(missing_ok=True)
            return False
        return True
    except Exception:
        hedef.unlink(missing_ok=True)
        return False


def pptx_gomulu_cikarici(pptx_yol: Path, gecici: Path):
    """Slayttaki tüm gömülü resimleri (resimler, nesne önizlemeleri, grup içindekiler) dosyadan çıkarır."""
    from pptx import Presentation
    from pptx.oxml.ns import qn
    sunu = Presentation(str(pptx_yol))

    def cikar(slayt_no0: int, klasor: Path, no: int):
        slayt = sunu.slides[slayt_no0]
        gorulen, kaydedilen = set(), []
        for el in slayt.shapes._spTree.iter(qn("a:blip")):
            rid = el.get(qn("r:embed"))
            if not rid:
                continue
            try:
                parca = slayt.part.related_part(rid)
            except KeyError:
                continue
            ozet = hashlib.sha1(parca.blob).hexdigest()
            if ozet in gorulen:
                continue
            gorulen.add(ozet)
            hedef = klasor / f"s{no:03d}_resim{len(kaydedilen) + 1}.png"
            if resim_kaydet(parca.blob, parca.partname.ext, hedef, gecici):
                kaydedilen.append(hedef.name)
        return kaydedilen
    return cikar


def pdf_gomulu_cikarici(belge):
    """PDF sayfasındaki gömülü resimleri (arka plan hariç) doğrudan çıkarır."""
    def cikar(sayfa_no: int, klasor: Path, no: int):
        sayfa = belge[sayfa_no]
        A = sayfa.rect.get_area()
        kaydedilen = []
        for bilgi in sayfa.get_images(full=True):
            xref = bilgi[0]
            try:
                kutular = sayfa.get_image_rects(xref)
                if not kutular or not (0.02 <= kutular[0].get_area() / A < 0.9):
                    continue
                pix = pymupdf.Pixmap(belge, xref)
                if pix.n - pix.alpha >= 4 or pix.alpha:
                    pix = pymupdf.Pixmap(pymupdf.csRGB, pix)
                hedef = klasor / f"s{no:03d}_resim{len(kaydedilen) + 1}.png"
                pix.save(hedef)
                oran = baskin_renk_orani(pix, (0, 0, 1, 1))
                if oran is None or oran >= BOS_ESIGI:
                    hedef.unlink(missing_ok=True)
                    continue
                kaydedilen.append(hedef.name)
            except Exception:
                continue
        return kaydedilen
    return cikar


# Bilgi notları: betiğin zaten hallettiği ya da içeriği etkilemeyen durumlar (kontrol gerektirmez).
BILGI_NOTLARI = ("metin görüntüden OCR ile okundu", "kelime okundu (OCR", "güvenilir okunamadı", "hatalı boşluk düzeltildi", "grafik (değerler sunudan okundu)", "gizli slayt", "başlık slayt çiziminden",
                 "ile yeniden çizildi", "görsel dosyadan çıkarıldı")


def kontrol_uyarilari(bayraklar):
    return [b for b in bayraklar if not any(n in b for n in BILGI_NOTLARI)]


def bilgi_notlari(bayraklar):
    return [b for b in bayraklar if any(n in b for n in BILGI_NOTLARI)]


GORSEL_ISARETLERI = ("görsel", "vektör şema", "şema/şekil", "şekil içi", "grafik", "hiç metin yok",
                     "PPTX ile PDF metni uyuşmuyor")


def gorsel_gerekli(yapi) -> bool:
    """Resim, şema, grafik içeren ya da hiç metni olmayan (taranmış) slaytların görüntüsü saklanır.
    Yalnızca yazıdan oluşan slaytların metni zaten slaytlar.md'de olduğu için görüntüsü atlanır."""
    if SAYFA_GORSELLERI == "hepsi":
        return True
    return any(b.startswith(GORSEL_ISARETLERI) for b in yapi["bayraklar"])


def slayt_md(no, yapi, kaynak_etiketi, gorsel_var=True, ek_resimler=()):
    parcalar = []
    baslik = yapi["baslik"] or "(başlık yok)"
    parcalar.append(f"## Slayt {no} — {baslik}")
    parcalar.append(f"<!-- kaynak: {kaynak_etiketi} -->")
    if gorsel_var:
        parcalar.append(f"![Slayt {no}](sayfalar/s{no:03d}.png)")
    for k, ad in enumerate(ek_resimler, 1):
        parcalar.append(f"![Slayt {no} — dosyadan çıkarılan resim {k}](sayfalar/{ad})")
    govde = []
    for tur, _, icerik in yapi["bloklar"]:
        if tur == "liste":
            satirlar = []
            for o in icerik:
                girinti = "  " * o["seviye"]
                if o["tur"] == "madde":
                    satirlar.append(f"{girinti}- {o['metin'].strip()}")
                else:
                    satirlar.append(f"{girinti}{md_kacis(o['metin'])}" if o["seviye"] == 0
                                    else f"{girinti}- {o['metin'].strip()}")
            # Paragraflar arasında boş satır, maddeler bitişik
            md = []
            for i, s in enumerate(satirlar):
                if i and (not s.lstrip().startswith("- ") or not satirlar[i - 1].lstrip().startswith("- ")):
                    md.append("")
                md.append(s)
            govde.append("\n".join(md))
        elif tur == "tablo" and icerik:
            genislik = max(len(r) for r in icerik)
            satirlar = [r + [""] * (genislik - len(r)) for r in icerik]
            md = ["| " + " | ".join(hucre(h) for h in satirlar[0]) + " |",
                  "|" + "---|" * genislik]
            md += ["| " + " | ".join(hucre(h) for h in r) + " |" for r in satirlar[1:]]
            govde.append("\n".join(md))
    if govde:
        parcalar.append("\n\n".join(govde))
    if BOS_UYARISI in yapi["bayraklar"]:
        parcalar.append("> ⛔ **Bu slaytın görseli boş çizildi.** PNG'de şema/resim yok; "
                        "içeriği asıl slayttan (PowerPoint veya Google Slaytlar) kontrol et.")
    if yapi["sekil_etiketleri"]:
        parcalar.append("> **Şekil içi etiketler** (yalnızca yazılar; oklar ve ilişkiler için görsele bak): " +
                        " · ".join(yapi["sekil_etiketleri"]))
    kontrol = [b for b in kontrol_uyarilari(yapi["bayraklar"]) if b != BOS_UYARISI]
    if kontrol:
        parcalar.append("> ⚠ **Kontrol:** " + "; ".join(kontrol))
    for ocr in yapi.get("ocr") or []:
        parca = [f"> **Görselden okunan yazı** (OCR, güven %{ocr['guven']:.0f} — otomatik okundu, "
                 f"şüphede kesin kaynak görseldir):"]
        if ocr["tablo"]:
            genislik = max(len(r) for r in ocr["tablo"])
            satirlar = [r + [""] * (genislik - len(r)) for r in ocr["tablo"]]
            parca.append(">")
            parca.append("> | " + " | ".join(hucre(h) for h in satirlar[0]) + " |")
            parca.append("> |" + "---|" * genislik)
            parca += ["> | " + " | ".join(hucre(h) for h in r) + " |" for r in satirlar[1:]]
        for cumle in ocr["etiketler"]:
            parca.append(">")
            parca.append("> " + cumle)
        parcalar.append("\n".join(parca))
    if yapi.get("not"):
        parcalar.append("**Konuşmacı notu:** " + re.sub(r"\s*\n\s*", " ", yapi["not"]))
    return "\n\n".join(parcalar)


def slayt_duz_metin(yapi):
    """Kapsama testi için çıktıdaki tüm metin (not hariç)."""
    parcalar = [yapi["baslik"]]
    for tur, _, icerik in yapi["bloklar"]:
        if tur == "liste":
            parcalar += [o["metin"] for o in icerik]
        else:
            parcalar += [h for r in icerik for h in r]
    parcalar += yapi["sekil_etiketleri"]
    return "\n".join(parcalar)


def paket_uret(deste_adi, kaynaklar, slaytlar, pdf_yol, cikis: Path, dpi, yontem, silinen_ustbilgi,
               pdf_metinleri=None, yedekler=None, gomulu_cikarici=None, rapor_yolu: Path = None):
    if cikis.exists():
        shutil.rmtree(cikis)
    (cikis / "sayfalar").mkdir(parents=True)

    # Görseller
    sayfa_sayisi = 0
    gorsel_eslesme = {}
    if pdf_yol:
        belge = pymupdf.open(pdf_yol)
        sayfa_sayisi = len(belge)
        gorunur = [i for i, s in enumerate(slaytlar) if not s.get("gizli")]
        if len(belge) == len(slaytlar):
            gorsel_eslesme = {i: i for i in range(len(slaytlar))}
        elif len(belge) == len(gorunur):
            gorsel_eslesme = {sl: sy for sy, sl in enumerate(gorunur)}

    # Markdown + ölçümler
    md_parcalari = [f"# {deste_adi} — slaytlar\n\n"
                    f"Kaynak: {', '.join(k.name for k in kaynaklar)} · yöntem: {yontem} · "
                    f"{len(slaytlar)} slayt. Metin dosyadan otomatik çıkarıldı; şemalar yorumlanmadı. "
                    f"Uyarılı slaytlar için `README.md` içindeki kalite tablosuna bak."]
    deste_sozlugu = Counter(k.lower() for sl in slaytlar
                            for k in re.findall(r"[A-Za-zÇĞİÖŞÜçğıöşü]+", slayt_duz_metin(sl)))
    if EK_SOZLUK:
        deste_sozlugu.update(EK_SOZLUK)
    olcumler = []
    for i, yapi in enumerate(slaytlar):
        no = i + 1
        pdf_sayfa = gorsel_eslesme.get(i)
        etiket = f"{kaynaklar[0].name} slayt {no}" + (f" / PDF s.{pdf_sayfa + 1}" if pdf_sayfa is not None else "")
        if yapi.get("gizli"):
            yapi["bayraklar"].append("gizli slayt")
        if yapi.get("ocr_sozluk_onar"):          # görüntü PDF: OCR'ın böldüğü/bitiştirdiği kelimeleri onar
            def duzelt(t):
                return ocr_metin_duzelt(ocr_kelime_onar(t, deste_sozlugu), deste_sozlugu)
            yapi["baslik"] = duzelt(yapi["baslik"])
            for tur, _, icerik in yapi["bloklar"]:
                if tur == "liste":
                    for o in icerik:
                        o["metin"] = duzelt(o["metin"])
                else:
                    icerik[:] = [[duzelt(h) for h in r] for r in icerik]

        # Kapsama: PDF modunda PyMuPDF düz çıkarımı; PPTX modunda PDF sayfasının metni (bağımsız kaynak)
        kaynak_metin = yapi.get("kaynak_metin")
        if kaynak_metin is None and pdf_sayfa is not None and pdf_metinleri:
            kaynak_metin = pdf_metinleri[pdf_sayfa]
        oran, eksik, fazla, toplam = (None, [], [], 0)
        if kaynak_metin is not None:
            oran, eksik, fazla, toplam = kapsama(kaynak_metin, slayt_duz_metin(yapi))
            # Düzeltilen boşluklar beklenen farklardır; ayrıca listelenir
            beklenen = set()
            for d in yapi["duzeltmeler"]:
                m = re.match(r"'(.*) (.*)' → '(.*)'", d)
                if m:
                    beklenen.update(kelimeler(m.group(1)) + kelimeler(m.group(2)) + kelimeler(m.group(3)))
            eksik = [k for k in eksik if k not in beklenen]
            fazla = [k for k in fazla if k not in beklenen]
            if yontem.startswith("pptx"):
                # İki farklı kaynak karşılaştırılıyor (sunu ↔ PDF); PDF bazı yazıları iki kez
                # çizebildiği için burada yalnızca çıktıda HİÇ geçmeyen kelimeler eksik sayılır.
                cikti_kelimeleri = set(kelimeler(slayt_duz_metin(yapi)))
                kaynak_kelimeleri = set(kelimeler(kaynak_metin))
                eksik = sorted(kaynak_kelimeleri - cikti_kelimeleri - beklenen)
                fazla = sorted(cikti_kelimeleri - kaynak_kelimeleri - beklenen)
                toplam = len(kaynak_kelimeleri)
            if toplam:
                oran = 1 - len(eksik) / toplam
            if yontem.startswith("pdf") and (eksik or fazla):
                yapi["bayraklar"].append("metin kapsama farkı")
            grafikli = any(b.startswith("grafik") for b in yapi["bayraklar"])
            if yontem.startswith("pptx") and toplam and oran < 0.95 and not grafikli:
                yapi["bayraklar"].append("PPTX ile PDF metni uyuşmuyor")
        doluluk, ek_resimler, cizen = [], [], None
        if yapi.get("kirpim") is not None:        # görüntü PDF: slaytın kendi kırpımı
            gorselli = gorsel_gerekli(yapi)
            if gorselli:
                yapi["kirpim"].save(cikis / "sayfalar" / f"s{no:03d}.png")
            etiket = f"{kaynaklar[0].name} {yapi.get('kaynak_konum', '')}"
        else:
            gorselli = pdf_sayfa is not None and gorsel_gerekli(yapi)
        if gorselli and yapi.get("kirpim") is None:
            png = cikis / "sayfalar" / f"s{no:03d}.png"
            belge[pdf_sayfa].get_pixmap(dpi=dpi).save(png)
            alanlar = yapi.get("gorsel_alanlari") or []
            if alanlar:
                doluluk = alanlari_olc(belge, pdf_sayfa, alanlar)
            # Güvence zinciri: 1) diğer çiziciler  2) resmi dosyadan çıkar  3) açık uyarı
            if bos_mu(doluluk):
                for arac, alt in (yedekler or []):
                    if len(alt) != len(belge):
                        continue
                    alt_doluluk = alanlari_olc(alt, pdf_sayfa, alanlar)
                    if not bos_mu(alt_doluluk):
                        alt[pdf_sayfa].get_pixmap(dpi=dpi).save(png)
                        doluluk, cizen = alt_doluluk, arac
                        yapi["bayraklar"].append(f"görsel {arac} ile yeniden çizildi (ilk çizim boştu)")
                        break
            if bos_mu(doluluk) and gomulu_cikarici:
                kaynak_no = pdf_sayfa if yontem.startswith("pdf") else i
                ek_resimler = gomulu_cikarici(kaynak_no, cikis / "sayfalar", no)
                if ek_resimler:
                    yapi["bayraklar"].append(f"görsel dosyadan çıkarıldı ({len(ek_resimler)} resim; "
                                             "slayt görüntüsündeki alan boştu)")
            if bos_mu(doluluk) and not ek_resimler:
                yapi["bayraklar"].append(BOS_UYARISI)
            elif ocr_hazir() and (alanlar or ek_resimler):
                cizim = next((alt for arac, alt in (yedekler or []) if arac == cizen), belge) if cizen else belge
                try:
                    okunan = slayt_ocr(cizim, pdf_sayfa, alanlar, [cikis / "sayfalar" / a for a in ek_resimler],
                                       deste_sozlugu)
                except Exception as e:
                    okunan = []
                    yapi["bayraklar"].append(f"OCR çalışmadı ({e})")
                def sozlukte_orani(o):
                    metin = " ".join([h for r in (o["tablo"] or []) for h in r] + o["etiketler"])
                    kel = [k.lower() for k in re.findall(r"[A-Za-zÇĞİÖŞÜçğıöşü]{3,}", metin)]
                    return sum(1 for k in kel if k in deste_sozlugu) / len(kel) if kel else 0

                buyuk_sozluk = len(deste_sozlugu) >= 1000
                guvenli = [o for o in okunan if o["guven"] >= OCR_GUVEN_ESIGI and o["kelime"] >= 2 and
                           (not buyuk_sozluk or sozlukte_orani(o) >= OCR_SOZLUK_ORANI)]
                if guvenli:
                    yapi["ocr"] = guvenli
                    yapi["bayraklar"] = ["görsel (yazısının okunabilen kısmı aşağıda; şema/ilişkiler için görsele bak)"
                                         if b.startswith("görsel (içindeki yazı çıkarılamaz") else b
                                         for b in yapi["bayraklar"]]
                    yapi["bayraklar"].append(
                        f"görselden {sum(o['kelime'] for o in guvenli)} kelime okundu (OCR, güven "
                        f"%{min(o['guven'] for o in guvenli):.0f})")
                elif okunan:
                    yapi["bayraklar"].append("görseldeki yazı güvenilir okunamadı (OCR güveni düşük)")
        md_parcalari.append(slayt_md(no, yapi, etiket, gorsel_var=gorselli, ek_resimler=ek_resimler))
        olcumler.append({
            "slayt": no, "gorsel": f"sayfalar/s{no:03d}.png" if gorselli else None,
            # gorsel_doluluk: görsel alanında baskın renk dışındaki piksel oranı (≈0 ise boş)
            "gorsel_doluluk": doluluk, "gorseli_cizen": cizen, "ek_resimler": ek_resimler,
            "ocr_kelime": sum(o["kelime"] for o in yapi.get("ocr") or []), "baslik": yapi["baslik"], "pdf_sayfa": None if pdf_sayfa is None else pdf_sayfa + 1,
            "kaynak_kelime": toplam, "kapsama": None if oran is None else round(oran, 4),
            "eksik_kelimeler": eksik[:15], "fazla_kelimeler": fazla[:15],
            "bosluk_duzeltmeleri": yapi["duzeltmeler"], "uyarilar": kontrol_uyarilari(yapi["bayraklar"]),
            "bilgiler": bilgi_notlari(yapi["bayraklar"]),
            "tablo": sum(1 for b in yapi["bloklar"] if b[0] == "tablo"),
            "not_var": bool(yapi.get("not")),
        })

    kaydedilen = len(list((cikis / "sayfalar").glob("*.png")))
    if not kaydedilen:
        (cikis / "sayfalar").rmdir()
    (cikis / "slaytlar.md").write_text("\n\n---\n\n".join(md_parcalari) + "\n", encoding="utf-8")

    # README: klasör açıklaması + slayt indeksi + kalite özeti
    kontrol = [o for o in olcumler if o["uyarilar"]]
    bos_gorseller = [o["slayt"] for o in olcumler if BOS_UYARISI in o["uyarilar"]]
    kurtarilan = [o["slayt"] for o in olcumler if o["gorseli_cizen"] or o["ek_resimler"]]
    gizli_sayisi = sum(1 for sl in slaytlar if sl.get("gizli"))
    gorselsiz = len(slaytlar) - len(gorsel_eslesme) - (gizli_sayisi if len(gorsel_eslesme) else 0)
    gorselsiz = sum(1 for i, sl in enumerate(slaytlar) if i not in gorsel_eslesme and not sl.get("gizli")
                    and gorsel_gerekli(sl)) if gorselsiz > 0 else 0
    r = [f"# {deste_adi}", "",
         f"- **Kaynak dosya(lar):** {', '.join(k.name for k in kaynaklar)}",
         f"- **Yöntem:** {yontem}",
         f"- **Slayt sayısı:** {len(slaytlar)} · **PDF sayfası:** {sayfa_sayisi}"
         + (f" · ⚠ {gorselsiz} görselli slaytın resmi üretilemedi (sayfa sayısı uyuşmadı)" if gorselsiz > 0 else "")
         + (f" · {gizli_sayisi} gizli slayt (PDF'te yer almaz, görseli yok)" if gizli_sayisi else ""),
         f"- **Görseli kaydedilen slayt:** {kaydedilen} (resim, şema veya grafik içerenler; "
         f"yalnızca yazı olan slaytların metni `slaytlar.md` içinde)",
         f"- **Kontrol edilmesi önerilen slayt:** {len(kontrol)} (`slaytlar.md` içinde `⚠ Kontrol` ya da `⛔` satırı olanlar; "
         f"ℹ ile gösterilen bilgi notları sayılmaz)",
         *([f"- 🛟 **İlk çizimi boş olup kurtarılan slayt:** {', '.join(map(str, kurtarilan))} "
            f"(başka çiziciyle yeniden çizildi ya da resim dosyadan çıkarıldı)"] if kurtarilan else []),
         *([f"- ⛔ **Görseli boş çizilen slayt:** {', '.join(map(str, bos_gorseller))} — bu PNG'lerde şema yok, "
            f"asıl slayta bak"] if bos_gorseller else []),
         "", "## Klasör", "",
         "| Dosya | İçerik |", "|---|---|",
         "| `slaytlar.md` | Tüm slaytların metni; her slayt `## Slayt N — Başlık` bölümü, görsel bağlantısı ve varsa uyarı ile |",
         "| `sayfalar/sNNN.png` | Yalnızca resim, şema veya grafik içeren slaytların görüntüsü (N = slayt numarası) |",
         "| `sayfalar/sNNN_resimK.png` | Slayt görüntüsü boş çıktığında dosyanın içinden doğrudan çıkarılan resim |",
         "| `README.md` | Bu dosya |",
         "", "## Nasıl kullanılır", "",
         "- Bir yapay zekâya ders notu hazırlatırken bütün klasörü değil, yalnızca ilgili slaytların "
         "`slaytlar.md` bölümlerini ver. Uyarılı slaytların PNG'sini de ekle.",
         "- `⚠` işaretli slaytlarda metin eksik olabilir veya şema yalnızca görselde vardır. Karar görsele göre verilir.",
         "- Kapsama, betiğin çıktısıyla PDF'in bağımsız metin çıkarımı arasındaki kelime örtüşmesidir; "
         "parantezdeki sayı ölçülen kelime sayısıdır. Görsele gömülü (resme dönüşmüş) yazıyı ölçmez: "
         "görselli slaytta az kelimeyle %100 görünmesi, yazının resimde olduğu anlamına gelebilir.",
         f"- Betiğin kendi ölçüm dosyası bu klasörde değil, `{JSON_KLASORU}/` klasöründedir; ders notu için gerekmez.",
         ]
    if silinen_ustbilgi:
        r += ["", "Her sayfada tekrar ettiği için çıkarılan üst/alt bilgi örnekleri: " +
              "; ".join(f"`{s}`" for s in silinen_ustbilgi[:3])]
    r += ["", "## Slayt indeksi", "", "| # | Başlık | Tablo | Not | Kapsama | Uyarılar |", "|---|---|---|---|---|---|"]
    for o in olcumler:
        kap = "—" if o["kapsama"] is None else f"%{o['kapsama'] * 100:.0f} ({o['kaynak_kelime']} kelime)"
        uyari = "; ".join(["⚠ " + u for u in o["uyarilar"]] + ["ℹ " + b for b in o["bilgiler"]]) or "—"
        if o["eksik_kelimeler"]:
            uyari += f" · eksik: {' '.join(o['eksik_kelimeler'][:6])}"
        if o["fazla_kelimeler"]:
            uyari += f" · fazla: {' '.join(o['fazla_kelimeler'][:6])}"
        r.append(f"| {o['slayt']} | {hucre(o['baslik']) or '—'} | {o['tablo'] or ''} | "
                 f"{'✓' if o['not_var'] else ''} | {kap} | {hucre(uyari)} |")
    duzeltmeler = [(o["slayt"], d) for o in olcumler for d in o["bosluk_duzeltmeleri"]]
    if duzeltmeler:
        r += ["", "## Otomatik boşluk düzeltmeleri", "",
              "PDF'te harflerin arasına normalden dar sahte boşluk yerleştirilmiş yerler birleştirildi. "
              "Hepsi burada listelenir; yanlış olan varsa görselle kontrol et.", ""]
        r += [f"- Slayt {no}: {d}" for no, d in duzeltmeler]
    (cikis / "README.md").write_text("\n".join(r) + "\n", encoding="utf-8")

    rapor_yolu = rapor_yolu or (cikis / "kalite_raporu.json")
    rapor_yolu.parent.mkdir(parents=True, exist_ok=True)
    rapor_yolu.write_text(json.dumps({
        "surum": SURUM, "deste": deste_adi, "kaynaklar": [k.name for k in kaynaklar],
        "kaynak_sha1": {k.name: sha1(k) for k in kaynaklar},
        "kaynak_md5": sorted(md5(k) for k in kaynaklar),   # Drive'ın md5Checksum'ı ile karşılaştırılır
        "yontem": yontem,
        "slaytlar": olcumler}, ensure_ascii=False, indent=2), encoding="utf-8")
    return olcumler


# ──────────────────────────────────────────────────────────────────────────────
# Ana akış
# ──────────────────────────────────────────────────────────────────────────────

def desteleri_bul(giris: Path):
    dosyalar = [p for p in giris.rglob("*") if p.is_file() and not p.name.startswith(("~$", "."))]
    gruplar = {}
    for p in dosyalar:
        uz = p.suffix.lower()
        if uz == ".pdf" or uz in SUNU_UZANTILARI:
            gruplar.setdefault((p.parent, p.stem), {})[uz] = p
    desteler = []
    for (klasor, ad), turler in sorted(gruplar.items(), key=lambda x: str(x[0])):
        sunu = next((turler[u] for u in (".pptx", ".ppsx", ".ppt", ".pps", ".odp") if u in turler), None)
        desteler.append({"ad": ad, "klasor": klasor, "sunu": sunu, "pdf": turler.get(".pdf")})
    return desteler


def guvenli_ad(ad: str) -> str:
    return re.sub(r"[^\w\-. ]+", "_", ad).strip() or "deste"


def gorselleri_birlestir(slaytlar, pdf_slaytlari):
    """Tanıma iki bağımsız kaynaktan: sunu yapısı + çizilmiş PDF'teki resimler.
    Sunu 'resim yok' dese bile PDF sayfasında resim varsa slayt resimli sayılır."""
    gorunur = [i for i, s in enumerate(slaytlar) if not s.get("gizli")]
    if len(pdf_slaytlari) == len(slaytlar):
        eslesme = list(range(len(slaytlar)))
    elif len(pdf_slaytlari) == len(gorunur):
        eslesme = gorunur
    else:
        return
    for sayfa_no, slayt_no in enumerate(eslesme):
        pdf_alanlari = pdf_slaytlari[sayfa_no].get("gorsel_alanlari") or []
        yapi = slaytlar[slayt_no]
        if pdf_alanlari and not any(b.startswith(GORSEL_ISARETLERI) for b in yapi["bayraklar"]):
            yapi["bayraklar"].append("görsel (çizilmiş PDF'te resim var, sunu yapısında görünmüyor)")
        for kutu in pdf_alanlari:
            alan_ekle_tekil(yapi.setdefault("gorsel_alanlari", []), kutu)


def rapor_yolu_bul(cikis_kok: Path, goreli: Path, ad: str) -> Path:
    return cikis_kok / JSON_KLASORU / goreli / f"{guvenli_ad(ad)}.json"


def basliklari_tamamla(slaytlar, pdf_slaytlari):
    """Dönüştürme sırasında başlık yer tutucusu kaybolmuşsa başlığı çizimden (PDF'teki en büyük
    yazı) alır ve aynı metni gövdeden çıkarır. İki bağımsız kaynaktan biri başlığı veriyorsa yeter."""
    gorunur = [i for i, s in enumerate(slaytlar) if not s.get("gizli")]
    if len(pdf_slaytlari) == len(slaytlar):
        eslesme = list(range(len(slaytlar)))
    elif len(pdf_slaytlari) == len(gorunur):
        eslesme = gorunur
    else:
        return

    def norm(t):
        return " ".join(kelimeler(t))
    for sayfa_no, slayt_no in enumerate(eslesme):
        yapi, pdf_baslik = slaytlar[slayt_no], pdf_slaytlari[sayfa_no].get("baslik", "")
        if yapi["baslik"] or not pdf_baslik:
            continue
        yapi["baslik"] = pdf_baslik
        hedef = norm(pdf_baslik)
        for blok in yapi["bloklar"]:
            if blok[0] == "liste":
                blok[2][:] = [o for o in blok[2] if norm(o["metin"]) != hedef]
        yapi["bloklar"] = [b for b in yapi["bloklar"] if b[0] != "liste" or b[2]]
        yapi["bayraklar"] = [b for b in yapi["bayraklar"] if b != "başlık yer tutucusu yok"]
        yapi["bayraklar"].append("başlık slayt çiziminden alındı (sunuda başlık alanı yok)")


def deste_isle(deste, giris: Path, cikis_kok: Path, dpi, yeniden):
    goreli = deste["klasor"].relative_to(giris)
    cikis = cikis_kok / goreli / guvenli_ad(deste["ad"])
    kaynaklar = [p for p in (deste["sunu"], deste["pdf"]) if p]
    rapor = rapor_yolu_bul(cikis_kok, goreli, deste["ad"])
    if rapor.exists() and cikis.exists() and not yeniden:
        try:
            eski = json.loads(rapor.read_text(encoding="utf-8"))
            if eski.get("surum") == SURUM and eski.get("kaynak_sha1") == {k.name: sha1(k) for k in kaynaklar}:
                return "atlandı (değişmemiş)", eski["slaytlar"], cikis
        except Exception:
            pass

    with tempfile.TemporaryDirectory() as gecici:
        gecici = Path(gecici)
        silinen = []
        if deste["sunu"]:
            sunu = deste["sunu"]
            metin_notu = ""
            if sunu.suffix.lower() != ".pptx":
                sunu, arac = sunu_cevir(sunu, "pptx", gecici)
                metin_notu = f", {deste['sunu'].suffix.lower()} → {arac} ile çevrildi"
            slaytlar = pptx_isle(sunu)
            gorunur = sum(1 for sl in slaytlar if not sl.get("gizli"))
            pdf_uyumlu = deste["pdf"] and len(pymupdf.open(deste["pdf"])) in (len(slaytlar), gorunur)
            arac = None
            if pdf_uyumlu:
                pdf, yontem = deste["pdf"], f"pptx (metin{metin_notu}) + pdf (görsel)"
            else:
                pdf, arac = sunu_cevir(deste["sunu"], "pdf", gecici)
                yontem = f"pptx (metin{metin_notu}) + {arac} (görsel)"
                if deste["pdf"]:
                    yontem += " — aynı adlı PDF'in sayfa sayısı uymadığı için kullanılmadı"
            pdf_slaytlari, silinen = pdf_isle(pdf)
            pdf_metinleri = [s["kaynak_metin"] for s in pdf_slaytlari]
            gorselleri_birlestir(slaytlar, pdf_slaytlari)
            basliklari_tamamla(slaytlar, pdf_slaytlari)
            kullanilan = "pdf" if pdf_uyumlu else arac
            adaylar = []
            if HARICI_DONUSTURUCU and kullanilan != "Google Slaytlar":
                adaylar.append(("Google Slaytlar", lambda: HARICI_DONUSTURUCU(deste["sunu"], "pdf", gecici / "g")))
            if kullanilan != "LibreOffice":
                adaylar.append(("LibreOffice", lambda: soffice_cevir(deste["sunu"], "pdf", gecici / "lo")))
            yedekler = YedekCizimler(adaylar)
            gomulu = pptx_gomulu_cikarici(sunu, gecici / "resim")
        else:
            pdf, pdf_metinleri, yedekler = deste["pdf"], None, None
            if goruntu_pdf_mi(pdf):
                if not ocr_hazir():
                    raise RuntimeError("Bu PDF yalnızca resimlerden oluşuyor (yazı katmanı yok); okumak için OCR "
                                       "(Tesseract) gerekli ama kurulu değil.")
                yontem = "görüntü PDF (yazı katmanı yok; metin OCR ile okundu)"
                slaytlar, silinen, gomulu = goruntu_pdf_isle(pdf), [], None
                for yapi in slaytlar:
                    yapi["ocr_sozluk_onar"] = True
            else:
                yontem = "pdf (metin ve görsel)"
                slaytlar, silinen = pdf_isle(pdf)
                gomulu = pdf_gomulu_cikarici(pymupdf.open(pdf))
        olcumler = paket_uret(deste["ad"], kaynaklar, slaytlar, pdf, cikis, dpi, yontem, silinen, pdf_metinleri,
                              yedekler, gomulu, rapor)
    return yontem, olcumler, cikis


def main():
    ap = argparse.ArgumentParser(description="Slaytları Markdown + görsel paketlerine çevirir.")
    ap.add_argument("giris", type=Path)
    ap.add_argument("cikis", type=Path)
    ap.add_argument("--dpi", type=int, default=110)
    ap.add_argument("--yeniden", action="store_true")
    ap.add_argument("--sozluk", type=Path, default=None,
                    help="Kitap Markdown/TXT klasörü: bölünmüş kelimeleri onarırken ek sözlük olarak kullanılır")
    a = ap.parse_args()

    global EK_SOZLUK
    if a.sozluk:
        EK_SOZLUK = Counter()
        for f in list(a.sozluk.rglob("*.md")) + list(a.sozluk.rglob("*.txt")):
            EK_SOZLUK.update(k.lower() for k in re.findall(r"[A-Za-zÇĞİÖŞÜçğıöşü]+", f.read_text(errors="ignore")))
        print(f"Ek sözlük: {len(EK_SOZLUK)} farklı kelime ({a.sozluk})")

    giris, cikis = a.giris.resolve(), a.cikis.resolve()
    if OCR and not ocr_hazir():
        print("ℹ OCR atlanıyor: Tesseract kurulu değil (macOS: brew install tesseract tesseract-lang; "
              "ayrıca pip install pytesseract). Görsellerdeki yazı md'ye eklenmeyecek.")
    desteler = desteleri_bul(giris)
    if not desteler:
        sys.exit(f"'{giris}' içinde PDF/PPT/PPTX bulunamadı.")
    cikis.mkdir(parents=True, exist_ok=True)
    indeks = ["# Slayt paketleri — indeks", "",
              "| Deste | Yöntem | Slayt | Kontrol önerilen | Klasör |", "|---|---|---|---|---|"]
    for d in desteler:
        ad = f"{d['klasor'].relative_to(giris) / d['ad']}"
        try:
            yontem, olcumler, klasor = deste_isle(d, giris, cikis, a.dpi, a.yeniden)
            kontrol = sum(1 for o in olcumler if o["uyarilar"])
            print(f"✓ {ad}: {len(olcumler)} slayt, {kontrol} kontrol önerilen — {yontem}")
            bos = [o["slayt"] for o in olcumler if BOS_UYARISI in o.get("uyarilar", [])]
            if bos:
                print(f"  ⛔ Görseli BOŞ çizilen slayt: {', '.join(map(str, bos))} (README'ye bak)")
            goreli = klasor.relative_to(cikis).as_posix()
            indeks.append(f"| {hucre(ad)} | {yontem} | {len(olcumler)} | {kontrol} | [{goreli}]({goreli}/README.md) |")
        except Exception as e:
            print(f"✗ {ad}: {e}")
            indeks.append(f"| {hucre(ad)} | HATA | — | — | {hucre(str(e))} |")
    (cikis / "INDEKS.md").write_text("\n".join(indeks) + "\n", encoding="utf-8")
    print(f"\nBitti. Genel indeks: {cikis / 'INDEKS.md'}")


# ──────────────────────────────────────────────────────────────────────────────
# Colab akışı — ses ve kitap betikleriyle AYNI klasör linki (drive.mount yok)
# ──────────────────────────────────────────────────────────────────────────────
# Klasör adlarında büyük/küçük harf, Türkçe karakter, boşluk ve alt çizgi fark etmez.
KAYNAK_KLASORU = "slaytlar"          # yalnızca bunun altı OKUNUR
CIKTI_KLASORU = "Slayt_Paketleri"    # yalnızca buraya YAZILIR (Slaytlar'ın kardeşi)
KITAP_KLASORU = "kitap"              # varsa buradaki *_Bolumler/*.md sözlük olarak kullanılır
DPI = 110
YENIDEN = False    # True: paketi zaten olan desteleri de yeniden işle


def _anahtar(ad: str) -> str:
    ad = ad.translate(str.maketrans("ıİşŞğĞüÜöÖçÇ", "iIsSgGuUoOcC")).lower()
    return re.sub(r"[^a-z0-9]", "", ad)


def _drive_klasor(servis, ust, ad):
    tip = "application/vnd.google-apps.folder"
    kac = ad.replace("\\", "\\\\").replace("'", "\\'")
    var = servis.files().list(
        q=f"'{ust}' in parents and name = '{kac}' and mimeType = '{tip}' and trashed = false",
        fields="files(id)", supportsAllDrives=True, includeItemsFromAllDrives=True).execute()["files"]
    if var:
        return var[0]["id"]
    return servis.files().create(body={"name": ad, "mimeType": tip, "parents": [ust]},
                                 fields="id", supportsAllDrives=True).execute()["id"]


def _drive_cocuklar(servis, ust):
    sonuc, sayfa = {}, None
    while True:
        r = servis.files().list(q=f"'{ust}' in parents and trashed = false",
                                fields="nextPageToken, files(id, name, mimeType)", pageSize=1000,
                                pageToken=sayfa, supportsAllDrives=True,
                                includeItemsFromAllDrives=True).execute()
        for f in r["files"]:
            sonuc[f["name"]] = f
        sayfa = r.get("nextPageToken")
        if not sayfa:
            return sonuc


def _agaci_yukle(servis, ust_id, yerel: Path):
    """Yerel klasörü Drive'daki ust_id klasörüne birebir aynalar.
    Yalnızca bu (deste) klasörünün içindeki eski/fazla dosyaları çöpe atar."""
    import mimetypes
    from googleapiclient.http import MediaFileUpload
    uzak = _drive_cocuklar(servis, ust_id)
    yerel_adlar = set()
    for p in sorted(yerel.iterdir()):
        yerel_adlar.add(p.name)
        if p.is_dir():
            alt = uzak.get(p.name)
            alt_id = alt["id"] if alt and alt["mimeType"].endswith("folder") else _drive_klasor(servis, ust_id, p.name)
            _agaci_yukle(servis, alt_id, p)
            continue
        medya = MediaFileUpload(str(p), mimetype=mimetypes.guess_type(p.name)[0] or "application/octet-stream")
        if p.name in uzak and not uzak[p.name]["mimeType"].endswith("folder"):
            servis.files().update(fileId=uzak[p.name]["id"], media_body=medya, supportsAllDrives=True).execute()
        else:
            servis.files().create(body={"name": p.name, "parents": [ust_id]}, media_body=medya,
                                  fields="id", supportsAllDrives=True).execute()
    for ad, f in uzak.items():
        if ad not in yerel_adlar:
            servis.files().update(fileId=f["id"], body={"trashed": True}, supportsAllDrives=True).execute()


def _dosya_yukle(servis, ust_id, yerel: Path):
    """Tek dosyayı klasöre yükler; aynı adlı dosya varsa üzerine yazar."""
    from googleapiclient.http import MediaFileUpload
    medya = MediaFileUpload(str(yerel), mimetype="application/json")
    var = _drive_cocuklar(servis, ust_id).get(yerel.name)
    if var and not var["mimeType"].endswith("folder"):
        servis.files().update(fileId=var["id"], media_body=medya, supportsAllDrives=True).execute()
    else:
        servis.files().create(body={"name": yerel.name, "parents": [ust_id]}, media_body=medya,
                              fields="id", supportsAllDrives=True).execute()


def _indir(servis, dosya_id, hedef, gdown_modulu, deneme=3):
    """Dosyayı indirir; başarısızsa None döner (asla çalışmayı durdurmaz).
    Drive izni varsa kimlikli Drive API kullanılır: herkese açık link yolundaki
    'çok fazla erişim' kısıtlamasına takılmaz. İzin yoksa gdown yedek olarak kullanılır."""
    import time
    hedef = Path(hedef)
    hedef.parent.mkdir(parents=True, exist_ok=True)
    for i in range(deneme):
        try:
            if servis:
                from googleapiclient.http import MediaIoBaseDownload
                istek = servis.files().get_media(fileId=dosya_id, supportsAllDrives=True)
                with open(hedef, "wb") as f:
                    indirici = MediaIoBaseDownload(f, istek, chunksize=16 * 1024 * 1024)
                    bitti = False
                    while not bitti:
                        _, bitti = indirici.next_chunk()
                return hedef
            if gdown_modulu.download(id=dosya_id, output=str(hedef), quiet=True):
                return hedef
        except Exception:
            time.sleep(2 * (i + 1))
    hedef.unlink(missing_ok=True)
    return None


def _yukari_bul(servis, dosya_id, hedef_anahtar):
    """Dosyadan yukarı çıkarak adı hedef_anahtar olan klasörün ÜST klasör id'sini bulur."""
    simdiki = servis.files().get(fileId=dosya_id, fields="parents", supportsAllDrives=True).execute()["parents"][0]
    for _ in range(30):
        bilgi = servis.files().get(fileId=simdiki, fields="name, parents", supportsAllDrives=True).execute()
        if _anahtar(bilgi["name"]) == hedef_anahtar:
            return bilgi["parents"][0]
        simdiki = bilgi["parents"][0]
    raise RuntimeError("Slaytlar klasörü Drive'da bulunamadı.")


def google_donusturucu(servis):
    """Sunuyu Drive'a geçici Google Slaytlar dosyası olarak yükler, PDF/PPTX olarak dışa aktarır.
    Geçici dosya her desteden sonra kalıcı olarak silinir (temizle)."""
    import mimetypes
    from googleapiclient.http import MediaFileUpload
    onbellek = {}
    turler = {"pdf": "application/pdf",
              "pptx": "application/vnd.openxmlformats-officedocument.presentationml.presentation"}
    kaynak_turleri = {".ppt": "application/vnd.ms-powerpoint", ".pps": "application/vnd.ms-powerpoint",
                      ".pptx": turler["pptx"], ".odp": "application/vnd.oasis.opendocument.presentation"}

    def cevir(kaynak: Path, hedef_tur: str, cikis: Path) -> Path:
        anahtar = str(kaynak)
        if anahtar not in onbellek:
            medya = MediaFileUpload(str(kaynak), resumable=True,
                                    mimetype=kaynak_turleri.get(kaynak.suffix.lower())
                                    or mimetypes.guess_type(kaynak.name)[0] or "application/octet-stream")
            onbellek[anahtar] = servis.files().create(
                body={"name": f"~gecici_slayt_paketle_{kaynak.name}",
                      "mimeType": "application/vnd.google-apps.presentation"},
                media_body=medya, fields="id").execute()["id"]
        veri = servis.files().export(fileId=onbellek[anahtar], mimeType=turler[hedef_tur]).execute()
        cikis.mkdir(parents=True, exist_ok=True)
        hedef = cikis / f"{kaynak.stem}.{hedef_tur}"
        hedef.write_bytes(veri)
        return hedef

    def temizle():
        for fid in onbellek.values():
            try:
                servis.files().delete(fileId=fid).execute()
            except Exception:
                pass
        onbellek.clear()

    cevir.temizle = temizle
    return cevir


def _libreoffice_kur():
    print("   🔧 LibreOffice kuruluyor (yedek dönüştürücü, ~1-2 dk)...")
    subprocess.run("apt-get -qq update > /dev/null; apt-get -qq install -y libreoffice-impress > /dev/null",
                   shell=True)


def colab_calistir():
    global EK_SOZLUK, HARICI_DONUSTURUCU, LIBREOFFICE_KURUCU
    import logging
    logging.getLogger("google_auth_httplib2").setLevel(logging.ERROR)   # zararsız zaman aşımı uyarıları
    subprocess.check_call([sys.executable, "-m", "pip", "install", "-q", "-U", "gdown", "python-pptx"])
    import gdown

    print("=" * 60)
    print("🖼  SLAYTLARI PAKETLE (Markdown + görsel)")
    print("=" * 60)
    print("Ses ve kitap betiklerinde kullandığın AYNI klasör linkini yapıştır.")
    print("(Yalnızca 'Slaytlar' klasörü okunur, yalnızca 'Slayt_Paketleri' klasörüne yazılır.)\n")
    link = input("🔗 Drive klasör linki: ").strip()
    if "folders/" in link:
        klasor_id = link.split("folders/")[1].split("?")[0].split("/")[0]
    elif "id=" in link:
        klasor_id = link.split("id=")[1].split("&")[0]
    else:
        klasor_id = link

    is_kok = Path("/content/slayt_is")
    shutil.rmtree(is_kok, ignore_errors=True)
    is_kok.mkdir(parents=True)

    print("\n🔍 Klasör listeleniyor...")
    try:
        liste = gdown.download_folder(id=klasor_id, output=str(is_kok / "liste") + "/",
                                      skip_download=True, quiet=True) or []
    except Exception as e:
        raise SystemExit(f"❌ Klasör okunamadı: {e}\n"
                         "   Linkin 'Bağlantıya sahip olan herkes' olarak paylaşıldığından emin ol.")
    yollar = {f.path.replace("\\", "/"): f for f in liste}

    # Slaytlar/ altındaki kaynaklar -> (slaytlar_yolu, göreli klasör, ad) grupları
    desteler, sozluk_dosyalari = {}, []
    cikti_anahtari = _anahtar(CIKTI_KLASORU)
    for yol, f in yollar.items():
        parcalar = yol.split("/")
        anahtarlar = [_anahtar(p) for p in parcalar[:-1]]
        uz = Path(parcalar[-1]).suffix.lower()
        if parcalar[-1].startswith(("~$", ".")):
            continue
        if KITAP_KLASORU in anahtarlar and yol.lower().endswith(".md") and \
                len(parcalar) >= 2 and _anahtar(parcalar[-2]).endswith("bolumler"):
            sozluk_dosyalari.append(f)
            continue
        if cikti_anahtari in anahtarlar or KAYNAK_KLASORU not in anahtarlar:
            continue
        if uz != ".pdf" and uz not in SUNU_UZANTILARI:
            continue
        i = anahtarlar.index(KAYNAK_KLASORU)
        slayt_kok = "/".join(parcalar[:i + 1])
        goreli = "/".join(parcalar[i + 1:-1])
        desteler.setdefault((slayt_kok, goreli, Path(parcalar[-1]).stem), {})[uz] = f

    if not desteler:
        raise SystemExit("\n⚠ 'Slaytlar' klasöründe PDF/PPT/PPTX bulunamadı.\n"
                         "   Beklenen yer: <Ders>/DersKaynaklari/Slaytlar/<dosya>")

    def paket_yolu(slayt_kok, goreli, ad):
        ust = slayt_kok.rsplit("/", 1)[0] if "/" in slayt_kok else ""
        return "/".join(x for x in (ust, CIKTI_KLASORU, goreli, guvenli_ad(ad)) if x)

    def json_yolu(slayt_kok, goreli, ad):
        ust = slayt_kok.rsplit("/", 1)[0] if "/" in slayt_kok else ""
        return "/".join(x for x in (ust, CIKTI_KLASORU, JSON_KLASORU, goreli, guvenli_ad(ad) + ".json") if x)

    servis = None
    print("\n☁️  Drive'a yükleme için izin isteniyor (istemezsen pencereyi kapat)...")
    try:
        from google.colab import auth
        from googleapiclient.discovery import build
        auth.authenticate_user()
        servis = build("drive", "v3", cache_discovery=False)
        print("✅ Drive izni alındı.")
    except Exception as e:
        print(f"⚠ Drive izni alınamadı ({e}). Sonuçlar zip olarak inecek.")

    def guncel_mi(anahtar, turler):
        """Paket var mı ve kaynak dosyalar paketlendikten sonra değişmemiş mi?"""
        rapor = yollar.get(json_yolu(*anahtar))
        if not rapor or paket_yolu(*anahtar) + "/slaytlar.md" not in yollar:
            return False
        if not servis:
            return True   # izin yoksa içerik karşılaştırılamaz; yalnızca varlığa bakılır
        try:
            hedef = is_kok / "raporlar" / f"{rapor.id}.json"
            hedef.parent.mkdir(exist_ok=True)
            if not _indir(servis, rapor.id, hedef, gdown):
                return False
            eski = json.loads(hedef.read_text(encoding="utf-8"))
            uzak = [servis.files().get(fileId=f.id, fields="md5Checksum", supportsAllDrives=True)
                    .execute().get("md5Checksum") for f in turler.values()]
            if None in uzak:          # Google Slaytlar gibi yerel biçimlerde md5 yok
                return eski.get("surum") == SURUM
            return eski.get("surum") == SURUM and eski.get("kaynak_md5") == sorted(uzak)
        except Exception:
            return False

    isler, atlanan = [], []
    for anahtar, turler in sorted(desteler.items()):
        (atlanan if not YENIDEN and guncel_mi(anahtar, turler) else isler).append((anahtar, turler))
    if not servis:
        print("ℹ Drive izni olmadığı için değişen slaytlar algılanamaz; yalnızca paketi hiç olmayanlar işlenir.")

    print(f"\n🖼  {len(desteler)} slayt destesi  |  ✅ {len(atlanan)} zaten paketlenmiş  |  🆕 {len(isler)} işlenecek")
    for (k, g, a), _ in atlanan:
        print(f"      ✅ {'/'.join(x for x in (g, a) if x)}")
    for (k, g, a), t in isler:
        print(f"      🆕 {'/'.join(x for x in (g, a) if x)}  ({', '.join(sorted(t))})")
    if not isler:
        raise SystemExit("\n✨ Paketlenecek yeni slayt yok. (Yeniden üretmek için YENIDEN = True yap "
                         "veya Slayt_Paketleri'ndeki o destenin klasörünü sil.)")

    if sozluk_dosyalari:
        EK_SOZLUK = Counter()
        sozluk_klasoru = is_kok / "sozluk"
        sozluk_klasoru.mkdir()
        inemeyen = 0
        for sira, f in enumerate(sozluk_dosyalari):
            hedef = sozluk_klasoru / f"{sira:03d}.md"
            if _indir(servis, f.id, hedef, gdown):
                EK_SOZLUK.update(k.lower() for k in re.findall(r"[A-Za-zÇĞİÖŞÜçğıöşü]+",
                                                                 hedef.read_text(errors="ignore")))
            else:
                inemeyen += 1
        print(f"📚 Kitap sözlüğü: {len(EK_SOZLUK)} farklı kelime ({len(sozluk_dosyalari) - inemeyen} bölüm dosyası"
              + (f"; {inemeyen} dosya indirilemedi, sözlüksüz devam ediliyor" if inemeyen else "") + ")")

    LIBREOFFICE_KURUCU = _libreoffice_kur
    if OCR:
        print("🔤 OCR (görsele dönüşmüş yazıyı okumak için) kuruluyor, ~20-40 sn...")
        subprocess.run("apt-get -qq update > /dev/null; apt-get -qq install -y tesseract-ocr tesseract-ocr-tur > /dev/null",
                       shell=True)
        subprocess.run([sys.executable, "-m", "pip", "install", "-q", "pytesseract"])
        _OCR_DURUM.clear()
        print("   ✅ OCR hazır." if ocr_hazir() else "   ⚠ OCR kurulamadı; görsellerdeki yazı okunmayacak.")
    if servis:
        HARICI_DONUSTURUCU = google_donusturucu(servis)
        if any(uz in SUNU_UZANTILARI for _, t in isler for uz in t):
            print("🔄 PPT/PPTX dosyaları Google Slaytlar ile çevrilecek (Drive'da geçici kopya açılıp silinir).")

    kaynak_yerel = is_kok / "kaynak"
    cikti_yerel = is_kok / "cikti"
    paket_kok_idleri = {}
    yuklenen, hatali, yuklenemeyen, bos_toplam = 0, 0, [], []
    for (slayt_kok, goreli, ad), turler in isler:
        etiket = "/".join(x for x in (goreli, ad) if x)
        print(f"\n⏳ {etiket}")
        yerel = {}
        try:
            giris = kaynak_yerel / slayt_kok
            yerel_klasor = giris / goreli
            yerel_klasor.mkdir(parents=True, exist_ok=True)
            for uz, f in turler.items():
                hedef = yerel_klasor / (ad + uz)
                if not _indir(servis, f.id, hedef, gdown):
                    raise RuntimeError(f"{ad + uz} indirilemedi (Drive erişimi; birkaç dakika sonra yeniden "
                                       f"çalıştırınca yalnızca eksik desteler işlenir).")
                yerel[uz] = hedef
            sunu = next((yerel[u] for u in (".pptx", ".ppsx", ".ppt", ".pps", ".odp") if u in yerel), None)
            deste = {"ad": ad, "klasor": yerel_klasor, "sunu": sunu, "pdf": yerel.get(".pdf")}
            cikis_kok = cikti_yerel / paket_yolu(slayt_kok, "", "x").rsplit("/", 1)[0]
            yontem, olcumler, paket = deste_isle(deste, giris, cikis_kok, DPI, True)
            rapor_yerel = rapor_yolu_bul(cikis_kok, Path(goreli), ad)
            kontrol = sum(1 for o in olcumler if o["uyarilar"])
            print(f"   ✓ {len(olcumler)} slayt, {kontrol} kontrol önerilen — {yontem}")
            bos = [o["slayt"] for o in olcumler if BOS_UYARISI in o["uyarilar"]]
            if bos:
                bos_toplam.append((etiket, bos))
                print(f"   ⛔ Görseli BOŞ çizilen slayt: {', '.join(map(str, bos))}")
            if not servis:
                yuklenemeyen += [paket, rapor_yerel]
            else:
                try:
                    ilk = next(iter(turler.values()))
                    if slayt_kok not in paket_kok_idleri:
                        ust = _yukari_bul(servis, ilk.id, KAYNAK_KLASORU)
                        paket_kok_idleri[slayt_kok] = _drive_klasor(servis, ust, CIKTI_KLASORU)
                    hedef_id = paket_kok_idleri[slayt_kok]
                    for parca in [p for p in goreli.split("/") if p] + [guvenli_ad(ad)]:
                        hedef_id = _drive_klasor(servis, hedef_id, parca)
                    _agaci_yukle(servis, hedef_id, paket)
                    json_id = paket_kok_idleri[slayt_kok]
                    for parca in [JSON_KLASORU] + [p for p in goreli.split("/") if p]:
                        json_id = _drive_klasor(servis, json_id, parca)
                    _dosya_yukle(servis, json_id, rapor_yerel)
                    yuklenen += 1
                    print(f"   ☁️  Drive: {CIKTI_KLASORU}/{etiket}/")
                except Exception as e:
                    yuklenemeyen += [paket, rapor_yerel]
                    print(f"   ⚠ Drive'a yüklenemedi ({e}) — paket zip'e eklenecek.")
        except Exception as e:
            hatali += 1
            print(f"   ❌ {e}")
        finally:
            for p in yerel.values():
                p.unlink(missing_ok=True)
            if HARICI_DONUSTURUCU:
                HARICI_DONUSTURUCU.temizle()

    # Genel INDEKS.md (önceden paketlenmiş desteler dahil)
    if servis and paket_kok_idleri:
      try:
        for slayt_kok, kok_id in paket_kok_idleri.items():
            satirlar = ["# Slayt paketleri — indeks", "",
                        "| Deste | Yöntem | Slayt | Kontrol önerilen | Klasör |", "|---|---|---|---|---|"]
            for (k, goreli, ad), _ in sorted(desteler.items()):
                if k != slayt_kok:
                    continue
                rapor = yollar.get(json_yolu(k, goreli, ad))
                yerel_rapor = cikti_yerel / json_yolu(k, goreli, ad)
                if not yerel_rapor.exists() and rapor:
                    yerel_rapor.parent.mkdir(parents=True, exist_ok=True)
                    _indir(servis, rapor.id, yerel_rapor, gdown)
                goreli_paket = "/".join(x for x in (goreli, guvenli_ad(ad)) if x)
                try:
                    r = json.loads(yerel_rapor.read_text(encoding="utf-8"))
                    kontrol = sum(1 for o in r["slaytlar"] if o["uyarilar"])
                    satirlar.append(f"| {hucre(goreli_paket)} | {r['yontem']} | {len(r['slaytlar'])} | "
                                    f"{kontrol} | [{goreli_paket}]({goreli_paket}/README.md) |")
                except Exception:
                    satirlar.append(f"| {hucre(goreli_paket)} | HATA / paket yok | — | — | — |")
            indeks = is_kok / "INDEKS.md"
            indeks.write_text("\n".join(satirlar) + "\n", encoding="utf-8")
            from googleapiclient.http import MediaFileUpload
            uzak = _drive_cocuklar(servis, kok_id).get("INDEKS.md")
            medya = MediaFileUpload(str(indeks), mimetype="text/markdown")
            if uzak:
                servis.files().update(fileId=uzak["id"], media_body=medya, supportsAllDrives=True).execute()
            else:
                servis.files().create(body={"name": "INDEKS.md", "parents": [kok_id]}, media_body=medya,
                                      fields="id", supportsAllDrives=True).execute()
      except Exception as e:
        print(f"⚠ INDEKS.md güncellenemedi ({e}).")

    print("\n" + "=" * 60)
    print(f"📊 {len(isler)} deste: ☁️ {yuklenen} Drive'a yüklendi · 📦 {sum(1 for y in yuklenemeyen if y.is_dir())} zip'e · ❌ {hatali} işlenemedi")
    if yuklenen:
        print(f"   Drive: DersKaynaklari/{CIKTI_KLASORU}/")
    if bos_toplam:
        print("⛔ Görseli boş çizilen slaytlar var (şema PNG'ye gelmedi):")
        for etiket, bos in bos_toplam:
            print(f"   • {etiket}: slayt {', '.join(map(str, bos))}")
        print("   Çözüm: aynı adla PowerPoint'ten alınmış bir PDF'i Slaytlar klasörüne koy; "
              "betik görselleri o PDF'ten alır.")
    if yuklenemeyen:
        zip_kok = is_kok / "zip"
        for yol in yuklenemeyen:
            if not yol.exists():
                continue
            hedef = zip_kok / yol.relative_to(cikti_yerel)
            if yol.is_dir():
                shutil.copytree(yol, hedef, dirs_exist_ok=True)
            else:
                hedef.parent.mkdir(parents=True, exist_ok=True)
                shutil.copy2(yol, hedef)
        zip_yolu = shutil.make_archive("/content/slayt_paketleri", "zip", zip_kok)
        try:
            from google.colab import files
            files.download(zip_yolu)
            print(f"⬇️  {Path(zip_yolu).name} indiriliyor (Drive'a yüklenemeyen paketler).")
        except Exception:
            print(f"⬇️  Sol paneldeki 📁'den {zip_yolu} dosyasını indir.")

if __name__ == "__main__":
    if "ipykernel" in sys.modules:   # Colab hücresine yapıştırılıp çalıştırıldı
        colab_calistir()
    else:
        main()